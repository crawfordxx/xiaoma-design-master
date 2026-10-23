"""生成「复杂模板」用的 pptx 底层工具：真实企业模板常见、python-pptx 没直接给的构造。
主题色 / 主题字体、母版和版式上的占位符与装饰、页面级背景和配色映射、渐变文字、描边文字、半透明、发光、
竖排和旋转文字、页码域、隐藏页、原生表格和图表。
"""
import os

from lxml import etree
from PIL import ImageFont
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
IMG_RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
THEME = {"accent1": MSO_THEME_COLOR.ACCENT_1, "accent2": MSO_THEME_COLOR.ACCENT_2,
         "accent3": MSO_THEME_COLOR.ACCENT_3, "accent4": MSO_THEME_COLOR.ACCENT_4,
         "tx1": MSO_THEME_COLOR.TEXT_1, "tx2": MSO_THEME_COLOR.TEXT_2, "bg1": MSO_THEME_COLOR.BACKGROUND_1,
         "bg2": MSO_THEME_COLOR.BACKGROUND_2}
ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def cjk_font(size, bold=False):
    for p in ("/System/Library/Fonts/PingFang.ttc", "/System/Library/Fonts/Hiragino Sans GB.ttc",
              "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"):
        if os.path.exists(p):
            return ImageFont.truetype(p, size, index=(8 if bold and p.endswith("PingFang.ttc") else 0))
    return ImageFont.load_default()


# ---------- 主题 ----------
def set_theme(prs, name, colors, major, minor):
    """改主题色板和主题字体。major / minor = (西文, 中文)"""
    part = next(r.target_part for r in prs.slide_master.part.rels.values() if r.reltype.endswith("/theme"))
    root = etree.fromstring(part.blob)
    root.set("name", name)
    cs = root.find(f".//{{{A}}}clrScheme")
    cs.set("name", name)
    for key, val in colors.items():
        el = cs.find(f"{{{A}}}{key}")
        for ch in list(el):
            el.remove(ch)
        etree.SubElement(el, f"{{{A}}}srgbClr", val=val)
    fs = root.find(f".//{{{A}}}fontScheme")
    fs.set("name", name)
    for kind, (latin, ea) in (("majorFont", major), ("minorFont", minor)):
        f = fs.find(f"{{{A}}}{kind}")
        f.find(f"{{{A}}}latin").set("typeface", latin)
        f.find(f"{{{A}}}ea").set("typeface", ea)
        for s in f.findall(f"{{{A}}}font"):
            if s.get("script") in ("Hans", "Hant"):
                s.set("typeface", ea)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def master_text_styles(prs, title, body):
    """母版默认文字样式：title=(字号, 主题色, 加粗)，body=[(字号, 主题色), …] 按级别"""
    tx = prs.slide_master._element.find(qn("p:txStyles"))

    def put(lvl, size, color, bold=None, font="+mn"):
        d = lvl.find(qn("a:defRPr"))
        d.set("sz", str(int(size * 100)))
        if bold is not None:
            d.set("b", "1" if bold else "0")
        for ch in list(d):
            d.remove(ch)
        f = etree.SubElement(d, qn("a:solidFill"))
        etree.SubElement(f, qn("a:schemeClr"), val=color)
        etree.SubElement(d, qn("a:latin"), typeface=f"{font}-lt")
        etree.SubElement(d, qn("a:ea"), typeface=f"{font}-ea")
    put(tx.find(qn("p:titleStyle")).find(qn("a:lvl1pPr")), *title, font="+mj")
    bs = tx.find(qn("p:bodyStyle"))
    for i, (size, color) in enumerate(body):
        put(bs.find(qn(f"a:lvl{i + 1}pPr")), size, color)


# ---------- 版式 ----------
def keep_layouts(prs, names):
    """只留下用到的版式并改名：names = {原序号: 新名字}"""
    master = prs.slide_master
    lst = master._element.find(qn("p:sldLayoutIdLst"))
    for i, (sid, layout) in enumerate(zip(list(lst), list(prs.slide_layouts))):
        if i in names:
            layout.name = names[i]
        else:
            master.part.drop_rel(sid.get(qn("r:id")))
            lst.remove(sid)
    return {n: lay for lay in prs.slide_layouts for n in [lay.name]}


def ph_of(container, ph_type=None, idx=None):
    for p in container.placeholders:
        if (idx is not None and p.placeholder_format.idx == idx) or \
                (ph_type is not None and p.placeholder_format.type == ph_type):
            return p
    return None


def style_ph(ph, x, y, w, h, prompt=None, size=None, color=None, bold=None, font=None, align=None, anchor=None,
             bullets=None):
    """设置版式 / 母版占位符的位置、提示文字和默认格式（写在 lstStyle 里，页面上的占位符继承它）"""
    ph.left, ph.top, ph.width, ph.height = Inches(x), Inches(y), Inches(w), Inches(h)
    tf = ph.text_frame
    if anchor:
        tf.vertical_anchor = ANCHOR[anchor]
    body = tf._txBody
    lst = body.find(qn("a:lstStyle"))
    for ch in list(lst):
        lst.remove(ch)
    lv = etree.SubElement(lst, qn("a:lvl1pPr"))
    if align:
        lv.set("algn", {"l": "l", "c": "ctr", "r": "r"}[align])
    if bullets is False:
        lv.set("marL", "0")
        lv.set("indent", "0")
        etree.SubElement(lv, qn("a:buNone"))
    d = etree.SubElement(lv, qn("a:defRPr"))
    if size:
        d.set("sz", str(int(size * 100)))
    if bold is not None:
        d.set("b", "1" if bold else "0")
    if color:
        f = etree.SubElement(d, qn("a:solidFill"))
        if color in THEME or color in ("dk1", "lt1", "dk2", "lt2"):
            etree.SubElement(f, qn("a:schemeClr"), val=color)
        else:
            etree.SubElement(f, qn("a:srgbClr"), val=color)
    if font:
        etree.SubElement(d, qn("a:latin"), typeface=f"{font}-lt" if font.startswith("+") else font)
        etree.SubElement(d, qn("a:ea"), typeface=f"{font}-ea" if font.startswith("+") else font)
    if prompt is not None:
        tf.text = prompt
    return ph


def drop_ph(container, *types_or_idx):
    for p in list(container.placeholders):
        f = p.placeholder_format
        if f.type in types_or_idx or f.idx in types_or_idx:
            p._element.getparent().remove(p._element)


def move_to(part_owner, slide, shapes, back=True):
    """把临时页上画好的形状搬到母版 / 版式上；back=True 放在最底层（占位符之下）"""
    tree = part_owner.shapes._spTree
    anchor = 2
    for s in shapes:
        el = s._element
        for blip in el.iter(qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            blip.set(qn("r:embed"), part_owner.part.relate_to(slide.part.related_part(rid), IMG_RT))
        if back:
            tree.insert(anchor, el)
            anchor += 1
        else:
            tree.append(el)


def layout_bg(layout_or_slide, color=None, theme=None):
    """页面 / 版式背景色（p:bg），可以是主题色"""
    cSld = layout_or_slide._element.find(qn("p:cSld"))
    old = cSld.find(qn("p:bg"))
    if old is not None:
        cSld.remove(old)
    bg = etree.Element(qn("p:bg"))
    pr = etree.SubElement(bg, qn("p:bgPr"))
    f = etree.SubElement(pr, qn("a:solidFill"))
    if theme:
        etree.SubElement(f, qn("a:schemeClr"), val=theme)
    else:
        etree.SubElement(f, qn("a:srgbClr"), val=color)
    etree.SubElement(pr, qn("a:effectLst"))
    cSld.insert(0, bg)


def dark_color_map(slide):
    """深色页：页面级配色映射，把文字 / 背景主题色对调（占位符的 tx1/tx2 自动变浅色）"""
    el = slide._element
    old = el.find(qn("p:clrMapOvr"))
    if old is not None:
        el.remove(old)
    ovr = etree.Element(qn("p:clrMapOvr"))
    etree.SubElement(ovr, qn("a:overrideClrMapping"), bg1="dk1", tx1="lt1", bg2="dk2", tx2="lt2",
                     accent1="accent1", accent2="accent2", accent3="accent3", accent4="accent4",
                     accent5="accent5", accent6="accent6", hlink="hlink", folHlink="folHlink")
    el.find(qn("p:cSld")).addnext(ovr)


def footer_phs(slide, master, number):
    """页面上放页脚和页码占位符（PowerPoint 只显示页面上有的那几个）：位置继承版式，文字取母版上的"""
    import copy
    tree = slide.shapes._spTree
    lay = {p.placeholder_format.type: p for p in slide.slide_layout.placeholders}
    for ph in master.placeholders:
        if ph.placeholder_format.type in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER):
            el = copy.deepcopy(ph._element)
            src = lay.get(ph.placeholder_format.type)
            if src is not None:  # 用版式上的 idx，页面占位符才能继承到位置
                el.find(".//" + qn("p:ph")).set("idx", str(src.placeholder_format.idx))
            used = [int(x.get("id")) for x in tree.iter(qn("p:cNvPr"))]
            el.find(".//" + qn("p:cNvPr")).set("id", str(max(used + [1]) + 1))  # 形状 id 在页面内要唯一
            sp_pr = el.find(qn("p:spPr"))
            for ch in list(sp_pr):  # 位置继承母版
                sp_pr.remove(ch)
            if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
                p = el.find(".//" + qn("a:p"))
                for ch in list(p):
                    if ch.tag != qn("a:pPr"):
                        p.remove(ch)
                fld = etree.SubElement(p, qn("a:fld"), id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
                t = etree.SubElement(fld, qn("a:t"))
                t.text = str(number)
            tree.append(el)


def hide(slide):
    slide._element.set("show", "0")


# ---------- 形状与文字 ----------
def fill(shape, color=None, theme=None, alpha=None, brightness=None):
    if color is None and theme is None:
        shape.fill.background()
        return shape
    shape.fill.solid()
    if theme:
        shape.fill.fore_color.theme_color = THEME[theme]
        if brightness:
            shape.fill.fore_color.brightness = brightness
    else:
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
    if alpha is not None:
        clr = shape.fill._xPr.find(qn("a:solidFill"))[0]
        etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 100000)))
    return shape


def outline(shape, color=None, width=0.75, theme=None, alpha=None):
    if color is None and theme is None:
        shape.line.fill.background()
        return shape
    if theme:
        shape.line.color.theme_color = THEME[theme]
    else:
        shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(width)
    if alpha is not None:
        clr = shape.line._ln.find(qn("a:solidFill"))[0]
        etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 100000)))
    return shape


def shape(shapes, kind, x, y, w, h, color=None, theme=None, alpha=None, line=None, line_w=0.75, rot=0, name=None,
          adj=None, brightness=None):
    s = shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(s, color, theme, alpha, brightness)
    outline(s, line, line_w)
    if adj is not None:
        for i, v in enumerate(adj if isinstance(adj, (list, tuple)) else [adj]):
            s.adjustments[i] = v
    if rot:
        s.rotation = rot
    if name:
        s.name = name
    s.text_frame.text = ""
    return s


def gradient(shape, stops, angle=0):
    """形状渐变填充：stops = [(位置 0-1, 颜色, 不透明度 0-1), …]"""
    spPr = shape._element.spPr
    for t in ("a:solidFill", "a:gradFill", "a:noFill"):
        old = spPr.find(qn(t))
        if old is not None:
            spPr.remove(old)
    g = etree.Element(qn("a:gradFill"), rotWithShape="1")
    gl = etree.SubElement(g, qn("a:gsLst"))
    for pos, c, a in stops:
        gs = etree.SubElement(gl, qn("a:gs"), pos=str(int(pos * 100000)))
        clr = etree.SubElement(gs, qn("a:srgbClr"), val=c)
        if a < 1:
            etree.SubElement(clr, qn("a:alpha"), val=str(int(a * 100000)))
    etree.SubElement(g, qn("a:lin"), ang=str(int(angle * 60000)), scaled="0")
    geom = spPr.find(qn("a:prstGeom"))
    (geom if geom is not None else spPr.find(qn("a:xfrm"))).addnext(g)
    return shape


def glow(shape, color, radius=8, alpha=0.5):
    spPr = shape._element.spPr
    eff = etree.Element(qn("a:effectLst"))
    gl = etree.SubElement(eff, qn("a:glow"), rad=str(int(radius * 12700)))
    clr = etree.SubElement(gl, qn("a:srgbClr"), val=color)
    etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 100000)))
    ln = spPr.find(qn("a:ln"))
    (ln if ln is not None else spPr[-1]).addnext(eff)


def text(shapes, x, y, w, h, paras, size=14, color=None, theme=None, bold=False, font="+mn", align="l",
         anchor="t", wrap=True, vert=None, rot=0, spacing=None, name=None, insets=None, autofit=False, spc=None):
    """文本框。paras：字符串（\\n 分段）或 [(文字, {size,color,theme,bold,font,align,grad,outline,spc}), …]
    font 写 '+mj' / '+mn' 表示用主题字体（真实模板最常见的写法）"""
    s = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = ANCHOR[anchor]
    bp = tf._txBody.find(qn("a:bodyPr"))
    if vert:
        bp.set("vert", vert)
    if insets is not None:
        for k, v in zip(("lIns", "tIns", "rIns", "bIns"), insets):
            bp.set(k, str(int(v * 914400)))
    if autofit:
        etree.SubElement(bp, qn("a:normAutofit"))
    if isinstance(paras, str):
        paras = [(p, {}) for p in paras.split("\n")]
    for i, (t, o) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[o.get("align", align)]
        if spacing or o.get("spacing"):
            p.line_spacing = o.get("spacing", spacing)
        if o.get("space_after"):
            p.space_after = Pt(o["space_after"])
        r = p.add_run()
        r.text = t
        style_run(r, size=o.get("size", size), color=o.get("color", color), theme=o.get("theme", theme),
                  bold=o.get("bold", bold), font=o.get("font", font), grad=o.get("grad"), stroke=o.get("outline"),
                  spc=o.get("spc", spc), alpha=o.get("alpha"))
    if rot:
        s.rotation = rot
    if name:
        s.name = name
    return s


def style_run(r, size=None, color=None, theme=None, bold=None, font=None, grad=None, stroke=None, spc=None,
              alpha=None):
    f = r.font
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    rPr = r._r.get_or_add_rPr()
    if spc:
        rPr.set("spc", str(int(spc * 100)))
    if stroke:  # 描边空心字
        ln = etree.SubElement(rPr, qn("a:ln"), w=str(int(stroke[1] * 12700)))
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=stroke[0])
        etree.SubElement(rPr, qn("a:noFill"))
    elif grad:
        g = etree.SubElement(rPr, qn("a:gradFill"))
        gl = etree.SubElement(g, qn("a:gsLst"))
        for pos, c in grad:
            gs = etree.SubElement(gl, qn("a:gs"), pos=str(int(pos * 100000)))
            etree.SubElement(gs, qn("a:srgbClr"), val=c)
        etree.SubElement(g, qn("a:lin"), ang="0", scaled="0")
    elif theme:
        f.color.theme_color = THEME[theme]
    elif color:
        f.color.rgb = RGBColor.from_string(color)
    if alpha is not None and (theme or color):
        clr = rPr.find(qn("a:solidFill"))[0]
        etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 100000)))
    if font:
        lt, ea = (f"{font}-lt", f"{font}-ea") if font.startswith("+") else (font, font)
        etree.SubElement(rPr, qn("a:latin"), typeface=lt)
        etree.SubElement(rPr, qn("a:ea"), typeface=ea)


def slide_number(shapes, x, y, w, h, n, size=9, color=None, theme=None, align="r"):
    """页码域：PowerPoint 里会自动显示当前页码"""
    s = text(shapes, x, y, w, h, str(n), size=size, color=color, theme=theme, align=align)
    p = s.text_frame.paragraphs[0]._p
    r = p.find(qn("a:r"))
    fld = etree.Element(qn("a:fld"), id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
    fld.append(r.find(qn("a:rPr")))
    t = etree.SubElement(fld, qn("a:t"))
    t.text = str(n)
    p.replace(r, fld)
    return s


def picture(shapes, path, x, y, w=None, h=None, crop=None, name=None):
    pic = shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if crop:
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
    if name:
        pic.name = name
    return pic


def group(shapes, items):
    """把已经画好的若干形状收进一个组合"""
    g = shapes.add_group_shape(items)
    return g


def table(shapes, x, y, w, h, rows, header_theme="accent1", font_size=11, band="F3F6FA", line="D5DCE6"):
    gt = shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    t = gt.table
    tblPr = t._tbl.tblPr
    for k in ("firstRow", "bandRow"):
        tblPr.set(k, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        tblPr.remove(sid)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = val
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            run = para.runs[0] if para.runs else para.add_run()
            style_run(run, size=font_size, bold=(r == 0), font="+mn",
                      theme="bg1" if r == 0 else "tx1")
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.theme_color = THEME[header_theme]
            else:
                cell.fill.fore_color.rgb = RGBColor.from_string(band if r % 2 == 0 else "FFFFFF")
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    return gt


def chart(shapes, kind, x, y, w, h, cats, series, colors, legend=True, labels=True, font_size=10, gap=80):
    data = CategoryChartData()
    data.categories = cats
    for name, vals in series:
        data.add_series(name, vals)
    gf = shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
    ch = gf.chart
    ch.has_legend = legend and len(series) > 1
    ch.font.size = Pt(font_size)
    for s, c in zip(ch.series, colors):
        s.format.fill.solid()
        if c in THEME:
            s.format.fill.fore_color.theme_color = THEME[c]
        else:
            s.format.fill.fore_color.rgb = RGBColor.from_string(c)
    if labels:
        ch.plots[0].has_data_labels = True
        ch.plots[0].data_labels.font.size = Pt(font_size)
    try:
        ch.plots[0].gap_width = gap
    except AttributeError:
        pass
    return gf
