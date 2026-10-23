"""复杂模板 2：科技发布会（虚构品牌「星曜科技」）。
考验点：
- Keynote 导出式结构：所有元素都直接放在页面上（空白版式），没有占位符
- 全幅照片背景、渐变填充、半透明玻璃卡片、发光、渐变文字、描边空心数字
- 竖排「目录」、旋转 90° 的英文、斜切的平行四边形面板
- 内容页右侧有出血的手机样机图，内容安全区不规则；深色页和浅色页混用
- 「左侧标题面板 + 右侧内容」的内容页（标题不在顶部）
"""
import math
import os
import random

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

import tpl_kit as K
from make_templates import OUT, rgb, save

W, H = 13.333, 7.5
PX = (1920, 1080)
INK, VIOLET, CYAN, MIST = "070914", "7C3AED", "22D3EE", "A5B4FC"
LAT, CN = "Montserrat", "思源黑体 CN"
GRAD = [(0, "FFFFFF"), (1, CYAN)]


# ---------- 美术 ----------
def space_bg(seed=2, floor=True):
    random.seed(seed)
    w, h = PX
    im = Image.new("RGBA", PX, rgb("05060F") + (255,))
    d = ImageDraw.Draw(im)
    for y in range(h):
        t = y / h
        d.line([(0, y), (w, y)], fill=(int(5 + 6 * t), int(6 + 8 * t), int(15 + 30 * t), 255))
    glow = Image.new("RGBA", PX, (0, 0, 0, 0))
    g = ImageDraw.Draw(glow)
    for cx, cy, r, c, a in ((1450, 300, 380, VIOLET, 150), (1700, 650, 300, CYAN, 90), (300, 900, 260, "4338CA", 90)):
        g.ellipse([cx - r, cy - r, cx + r, cy + r], fill=rgb(c) + (a,))
    im = Image.alpha_composite(im, glow.filter(ImageFilter.GaussianBlur(140)))
    d = ImageDraw.Draw(im)
    for _ in range(260):
        x, y, r = random.uniform(0, w), random.uniform(0, h * 0.7), random.choice((1, 1, 1, 2))
        d.ellipse([x - r, y - r, x + r, y + r], fill=(255, 255, 255, random.randint(60, 200)))
    if floor:
        hz = int(h * 0.72)
        for k in range(-24, 25):
            d.line([(w / 2 + k * 18, hz), (w / 2 + k * 160, h)], fill=rgb(CYAN) + (36,), width=1)
        for j in range(1, 12):
            y = hz + (h - hz) * (j / 12) ** 1.8
            d.line([(0, y), (w, y)], fill=rgb(CYAN) + (int(20 + j * 3),), width=1)
    return save(im, f"keynote-space{seed}{'' if floor else '-nf'}.jpg")


def dot_bg():
    w, h = PX
    im = Image.new("RGBA", PX, rgb("0A0D1C") + (255,))
    glow = Image.new("RGBA", PX, (0, 0, 0, 0))
    ImageDraw.Draw(glow).ellipse([1300, -300, 2200, 500], fill=rgb(VIOLET) + (110,))
    im = Image.alpha_composite(im, glow.filter(ImageFilter.GaussianBlur(160)))
    d = ImageDraw.Draw(im)
    for x in range(30, w, 48):
        for y in range(30, h, 48):
            d.point((x, y), fill=(255, 255, 255, 38))
    return save(im, "keynote-dots.jpg")


def orb():
    s = 1200
    im = Image.new("RGBA", (s, s), (0, 0, 0, 0))
    halo = Image.new("RGBA", (s, s), (0, 0, 0, 0))
    ImageDraw.Draw(halo).ellipse([150, 150, s - 150, s - 150], fill=rgb(VIOLET) + (120,))
    im = Image.alpha_composite(im, halo.filter(ImageFilter.GaussianBlur(90)))
    d = ImageDraw.Draw(im)
    for k in range(60):
        a0 = k * 6
        r = 330 + 40 * math.sin(k / 4)
        c = rgb(CYAN) if k % 3 else rgb(MIST)
        d.arc([s / 2 - r, s / 2 - r * 0.42, s / 2 + r, s / 2 + r * 0.42], a0, a0 + 150, fill=c + (90,), width=3)
    for r, a in ((250, 255), (180, 200), (110, 160)):
        d.ellipse([s / 2 - r, s / 2 - r, s / 2 + r, s / 2 + r], outline=(255, 255, 255, a // 2), width=2)
    core = Image.new("RGBA", (s, s), (0, 0, 0, 0))
    ImageDraw.Draw(core).ellipse([s / 2 - 120, s / 2 - 120, s / 2 + 120, s / 2 + 120], fill=(230, 245, 255, 230))
    im = Image.alpha_composite(im, core.filter(ImageFilter.GaussianBlur(40)))
    path = os.path.join(os.path.dirname(save(Image.new("RGB", (2, 2)), "_tmp.jpg")), "keynote-orb.png")
    im.save(path)
    return path


def phone():
    w, h = 760, 1500
    im = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    sh = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    ImageDraw.Draw(sh).rounded_rectangle([70, 60, w - 50, h - 20], 110, fill=rgb(VIOLET) + (120,))
    im = Image.alpha_composite(im, sh.filter(ImageFilter.GaussianBlur(40)))
    d = ImageDraw.Draw(im)
    d.rounded_rectangle([60, 40, w - 60, h - 40], 100, fill=(28, 30, 40, 255), outline=(120, 124, 140, 255), width=6)
    scr = Image.new("RGBA", (w - 160, h - 180))
    sd = ImageDraw.Draw(scr)
    for y in range(scr.height):
        t = y / scr.height
        sd.line([(0, y), (scr.width, y)], fill=(int(70 + 40 * t), int(40 + 120 * t), int(200 - 30 * t), 255))
    sd.rounded_rectangle([40, 140, scr.width - 40, 480], 40, fill=(255, 255, 255, 60))
    for k in range(4):
        sd.rounded_rectangle([40, 540 + k * 170, scr.width - 40, 680 + k * 170], 30, fill=(255, 255, 255, 40))
    sd.ellipse([scr.width / 2 - 80, 220, scr.width / 2 + 80, 380], fill=(255, 255, 255, 140))
    mask = Image.new("L", scr.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle([0, 0, scr.width, scr.height], 80, fill=255)
    im.paste(scr, (80, 90), mask)
    d.rounded_rectangle([w / 2 - 90, 70, w / 2 + 90, 110], 20, fill=(10, 10, 14, 255))
    path = os.path.join(os.path.dirname(save(Image.new("RGB", (2, 2)), "_tmp.jpg")), "keynote-phone.png")
    im.save(path)
    return path


# ---------- 小部件 ----------
def bg(s, path):
    return K.picture(s.shapes, path, 0, 0, W, H)


def logo(s, x, y, dark=True):
    star = K.shape(s.shapes, MSO_SHAPE.STAR_4_POINT, x, y, 0.34, 0.34, color=CYAN if dark else VIOLET)
    t = K.text(s.shapes, x + 0.42, y - 0.02, 1.8, 0.38, "XINGYAO", size=13, bold=True, font=LAT,
               color="FFFFFF" if dark else INK, anchor="m", spc=2)
    return K.group(s.shapes, [star, t])


def tag(s, dark=True):
    K.text(s.shapes, W - 3.2, 0.42, 2.6, 0.3, "XINGYAO LAUNCH 2027", size=9, font=LAT, align="r", spc=2,
           color="8B93B8" if dark else "8A8FA3")


def title(s, txt, dark=True, x=0.75, y=0.62, w=8.5, size=30, kicker=None):
    if kicker:
        K.text(s.shapes, x, y - 0.32, 4, 0.3, kicker, size=10, color=CYAN if dark else VIOLET, bold=True, font=LAT,
               spc=3)
    return K.text(s.shapes, x, y, w, 0.75, txt, size=size, bold=True, font=CN, color="FFFFFF" if dark else INK,
                  anchor="m")


def slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def make():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    art = {"space": space_bg(), "space2": space_bg(5, floor=False), "dots": dot_bg(), "orb": orb(), "phone": phone()}

    # 0 封面：全幅星空 + 光环产品图 + 渐变标题
    s = slide(prs)
    bg(s, art["space"])
    K.picture(s.shapes, art["orb"], 6.9, 0.2, 6.4)
    logo(s, 0.8, 0.7)
    K.text(s.shapes, 0.8, 2.15, 6, 0.35, "XINGYAO LAUNCH EVENT 2027", size=12, color=CYAN, font=LAT, spc=4, bold=True)
    K.text(s.shapes, 0.8, 2.6, 7.2, 1.4, [("产品发布会主标题", {"grad": GRAD})], size=54, bold=True, font=CN)
    K.text(s.shapes, 0.8, 4.0, 6.5, 0.5, "SUBTITLE · 在此输入副标题", size=18, color="C7CCE8", font=CN)
    line = K.shape(s.shapes, MSO_SHAPE.RECTANGLE, 0.82, 4.75, 1.6, 0.05)
    K.gradient(line, [(0, VIOLET, 1), (1, CYAN, 1)])
    K.text(s.shapes, 0.8, 5.0, 5, 0.4, "2027.03.18　·　深圳", size=13, color="8B93B8", font=LAT)

    # 1 目录：竖排「目录」+ 旋转英文 + 2×3 网格，序号是描边空心字
    s = slide(prs)
    K.layout_bg(s, color=INK)
    K.picture(s.shapes, art["space2"], 0, 0, W, H)
    K.text(s.shapes, 0.75, 1.0, 1.1, 2.6, "目录", size=44, bold=True, color="FFFFFF", font=CN, vert="eaVert")
    K.text(s.shapes, -0.55, 4.95, 3.0, 0.4, "CONTENTS", size=12, color=CYAN, font=LAT, spc=6, rot=270)
    K.shape(s.shapes, MSO_SHAPE.RECTANGLE, 2.25, 1.0, 0.03, 5.5, color="FFFFFF", alpha=0.18)
    for i in range(6):
        x, y = 2.9 + (i % 3) * 3.4, 1.25 + (i // 3) * 2.75
        n = K.text(s.shapes, x, y, 1.6, 0.95, [(f"{i + 1:02d}", {"outline": (CYAN, 1.25)})], size=48, bold=True,
                   font=LAT)
        t = K.text(s.shapes, x, y + 1.0, 3.0, 0.5, "目录标题", size=20, bold=True, color="FFFFFF", font=CN)
        e = K.text(s.shapes, x, y + 1.5, 3.0, 0.35, "TITLE HERE", size=10, color="8B93B8", font=LAT, spc=2)
        K.group(s.shapes, [n, t, e])

    # 2 章节 A：星空 + 斜切渐变面板 + 巨大描边序号 + 旋转英文
    s = slide(prs)
    bg(s, art["space"])
    panel = K.shape(s.shapes, MSO_SHAPE.PARALLELOGRAM, -1.2, 0, 8.6, H, adj=0.28)
    K.gradient(panel, [(0, VIOLET, 0.85), (1, "1E1B4B", 0.2)], angle=0)
    K.text(s.shapes, 0.9, 1.1, 4.5, 2.2, [("01", {"outline": ("FFFFFF", 1.5)})], size=150, bold=True, font=LAT)
    K.text(s.shapes, 10.6, 3.2, 3.6, 0.4, "CHAPTER ONE", size=12, color=CYAN, font=LAT, spc=6, rot=90)
    K.text(s.shapes, 0.95, 3.75, 6.5, 0.9, "章节标题", size=44, bold=True, color="FFFFFF", font=CN)
    K.text(s.shapes, 0.95, 4.65, 6, 0.45, "Chapter subtitle here", size=16, color="C7CCE8", font=LAT)

    # 3 章节 B（浅色）：渐变大圆 + 渐变序号
    s = slide(prs)
    circ = K.shape(s.shapes, MSO_SHAPE.OVAL, 7.4, -1.4, 7.8, 7.8)
    K.gradient(circ, [(0, CYAN, 0.9), (1, VIOLET, 0.9)], angle=45)
    K.text(s.shapes, 0.9, 1.6, 4, 1.6, [("02", {"grad": [(0, VIOLET), (1, CYAN)]})], size=110, bold=True, font=LAT)
    K.text(s.shapes, 0.95, 3.55, 6, 0.9, "章节标题", size=44, bold=True, color=INK, font=CN)
    K.text(s.shapes, 0.95, 4.45, 6, 0.45, "Chapter subtitle here", size=16, color="6B7280", font=LAT)
    logo(s, 0.95, 6.6, dark=False)

    # 4 内容（深色）+ 右下出血的手机样机
    s = slide(prs)
    bg(s, art["dots"])
    title(s, "页面标题", kicker="FEATURE 01")
    tag(s)
    K.picture(s.shapes, art["phone"], 9.0, 2.0, 3.6)
    K.text(s.shapes, 0.75, 1.9, 7.4, 3.0, "在此输入正文内容，介绍产品的一项核心能力。建议每页只讲一件事，用一句结论加两三个支撑点。",
           size=16, color="C7CCE8", font=CN, spacing=1.4)

    # 5 内容（深色，全宽）+ 渐变下划线
    s = slide(prs)
    bg(s, art["dots"])
    title(s, "页面标题", kicker="OVERVIEW")
    bar = K.shape(s.shapes, MSO_SHAPE.RECTANGLE, 0.77, 1.42, 1.2, 0.05)
    K.gradient(bar, [(0, VIOLET, 1), (1, CYAN, 1)])
    tag(s)

    # 6 内容（浅色）
    s = slide(prs)
    title(s, "页面标题", dark=False, kicker="DETAILS")
    K.shape(s.shapes, MSO_SHAPE.RECTANGLE, 0.77, 1.42, 1.2, 0.05, color=VIOLET)
    tag(s, dark=False)
    K.shape(s.shapes, MSO_SHAPE.RECTANGLE, 0, H - 0.12, W, 0.12, color=INK)

    # 7 示例：玻璃卡片 ×3
    s = slide(prs)
    bg(s, art["dots"])
    title(s, "功能亮点示例", kicker="HIGHLIGHTS")
    tag(s)
    for i in range(3):
        x = 0.75 + i * 4.0
        card = K.shape(s.shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.0, 3.7, 4.3, color="FFFFFF", alpha=0.06,
                       line="FFFFFF", adj=0.06)
        K.outline(card, "FFFFFF", 0.75, alpha=0.25)
        ic = K.shape(s.shapes, MSO_SHAPE.OVAL, x + 0.4, 2.45, 0.8, 0.8)
        K.gradient(ic, [(0, VIOLET, 1), (1, CYAN, 1)], angle=45)
        K.glow(ic, CYAN, 10, 0.35)
        t = K.text(s.shapes, x + 0.4, 3.55, 3.0, 0.5, "功能名称", size=20, bold=True, color="FFFFFF", font=CN)
        dsc = K.text(s.shapes, x + 0.4, 4.15, 2.9, 1.6, "在此输入功能描述文字，一两句话说清楚它解决了什么问题。",
                     size=12, color="A8AFCF", font=CN, spacing=1.35)
        K.group(s.shapes, [card, ic, t, dsc])

    # 8 示例：大数字参数
    s = slide(prs)
    bg(s, art["dots"])
    title(s, "核心参数示例", kicker="SPECS")
    tag(s)
    for i, (v, u) in enumerate((("XX", "GB"), ("XX", "h"), ("XX", "nm"), ("XX", "%"))):
        x = 0.75 + i * 3.1
        K.text(s.shapes, x, 2.6, 2.9, 1.3, [(v, {"grad": GRAD}), (u, {"size": 22, "color": CYAN})], size=72,
               bold=True, font=LAT)
        K.text(s.shapes, x, 4.4, 2.8, 0.4, "参数名称", size=14, color="FFFFFF", bold=True, font=CN)
        K.text(s.shapes, x, 4.85, 2.7, 0.7, "在此输入参数说明", size=11, color="8B93B8", font=CN)

    # 9 示例：参数对比表（浅色页，原生表格）
    s = slide(prs)
    title(s, "参数对比示例", dark=False, kicker="COMPARE")
    tag(s, dark=False)
    K.table(s.shapes, 0.75, 1.9, 11.8, 3.6, [["参数", "上一代", "本代", "提升"]] +
            [[f"参数名称 {i}", "XX", "XX", "+XX%"] for i in range(1, 6)], header_theme="tx1", font_size=13,
            band="F4F3FF", line="E5E7EB")
    K.shape(s.shapes, MSO_SHAPE.RECTANGLE, 0, H - 0.12, W, 0.12, color=INK)

    # 10 内容（左侧斜切标题面板，内容在右边）
    s = slide(prs)
    bg(s, art["dots"])
    panel = K.shape(s.shapes, MSO_SHAPE.PARALLELOGRAM, -1.0, 0, 6.2, H, adj=0.22)
    K.gradient(panel, [(0, "1E1B4B", 1), (1, VIOLET, 0.75)], angle=90)
    K.text(s.shapes, 0.75, 2.6, 3.6, 1.4, "页面标题", size=36, bold=True, color="FFFFFF", font=CN)
    K.text(s.shapes, 0.75, 3.9, 3.3, 1.0, "在此输入一句页面说明", size=13, color="C7CCE8", font=CN)
    tag(s)

    # 11 金句页
    s = slide(prs)
    bg(s, art["space2"])
    K.text(s.shapes, 1.2, 1.0, 2, 1.6, [("“", {"grad": [(0, VIOLET), (1, CYAN)]})], size=150, bold=True, font=LAT)
    K.text(s.shapes, 1.6, 2.7, 10, 1.6, "在这里输入一句有力量的话", size=40, bold=True, color="FFFFFF", font=CN)
    K.text(s.shapes, 1.6, 4.4, 6, 0.4, "—— 署名 / 职务", size=14, color="8B93B8", font=CN)

    # 12 结尾：渐变 THANK YOU + 二维码位
    s = slide(prs)
    bg(s, art["space"])
    K.text(s.shapes, 0.8, 2.2, 9, 1.4, [("THANK YOU", {"grad": GRAD})], size=80, bold=True, font=LAT)
    K.text(s.shapes, 0.85, 3.65, 6, 0.5, "感谢观看", size=22, color="FFFFFF", font=CN)
    qr = K.shape(s.shapes, MSO_SHAPE.ROUNDED_RECTANGLE, 10.6, 4.6, 1.7, 1.7, color="FFFFFF", adj=0.08)
    K.text(s.shapes, 10.6, 4.6, 1.7, 1.7, "二维码", size=12, color="9CA3AF", align="c", anchor="m")
    K.text(s.shapes, 10.2, 6.4, 2.5, 0.35, "扫码关注星曜科技", size=11, color="8B93B8", align="c", font=CN)
    logo(s, 0.85, 0.7)
    del qr

    path = os.path.join(OUT, "keynote-xingyao.pptx")
    prs.save(path)
    return path


if __name__ == "__main__":
    print(make())
