"""回归测试：python3 -m pytest tests/ -q（不需要 LibreOffice，渲染相关的用 --no-render）"""
import json
import os
import subprocess
import sys

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ENGINE = os.path.join(ROOT, "engine")
sys.path.insert(0, ENGINE)

from plan import outline, parse, unit_slides  # noqa: E402

TEMPLATES = {"finance-ruixin": [0, 1, 2, 3, 4, 5], "medical-kangning": [0, 1, 2, 3, 4],
             "education-qiming": [0, 1, 2, 3, 4], "manufacturing-huazhu": [0, 1, 2, 3, 4, 5],
             "tourism-yunqi": [0, 1, 2, 3, 4]}
EXPECT_ROLES = {"finance-ruixin": ["cover", "toc", "section", "content", "content", "ending"],
                "medical-kangning": ["cover", "toc", "section", "content", "ending"],
                "education-qiming": ["cover", "toc", "section", "content", "ending"],
                "manufacturing-huazhu": ["cover", "toc", "section", "content", "content", "ending"],
                "tourism-yunqi": ["cover", "toc", "section", "content", "ending"]}


COMPLEX_ROLES = {
    "energy-yuechuan": ["cover", "toc", "section"] + ["content"] * 9 + ["ending"],
    "keynote-xingyao": ["cover", "toc", "section", "section"] + ["content"] * 7 + ["quote", "ending"],
    "consulting-hengce": ["hidden", "cover", "toc", "section"] + ["content"] * 9 + ["ending"]}


def run(*args, cwd=ROOT):
    return subprocess.run([sys.executable, *args], cwd=cwd, capture_output=True, text=True)


@pytest.fixture(scope="module")
def profiles(tmp_path_factory):
    out = {}
    for name in TEMPLATES:
        d = tmp_path_factory.mktemp(name)
        r = run(os.path.join(ENGINE, "inspect_template.py"), f"examples/templates/{name}.pptx", "--out", str(d),
                "--no-render")
        assert r.returncode == 0, r.stderr
        out[name] = (d, json.load(open(d / "profile.json", encoding="utf-8")))
    return out


@pytest.fixture(scope="module")
def complex_profiles(tmp_path_factory):
    out = {}
    for name in COMPLEX_ROLES:
        d = tmp_path_factory.mktemp(name)
        r = run(os.path.join(ENGINE, "inspect_template.py"), f"examples/templates/{name}.pptx", "--out", str(d),
                "--no-render")
        assert r.returncode == 0, r.stderr
        out[name] = json.load(open(d / "profile.json", encoding="utf-8"))
    return out


@pytest.mark.parametrize("name", list(TEMPLATES))
def test_roles(profiles, name):
    _, prof = profiles[name]
    assert [s["role"] for s in prof["slides"]] == EXPECT_ROLES[name]


def test_brand_text_is_not_a_slot(profiles):
    _, prof = profiles["finance-ruixin"]
    cover = prof["slides"][0]
    brand = [t["key"] for t in cover["texts"] if t["text"] == "睿信资本"]
    assert brand and brand[0] not in sum([v if isinstance(v, list) else [v] for v in cover["slots"].values()], [])


def test_toc_items_found(profiles):
    for name in TEMPLATES:
        toc = profiles[name][1]["slides"][1]
        assert len(toc["slots"]["items"]) >= 4, name


def blocks(md):
    return outline(parse(md), "fallback")


def test_bom_title(profiles, tmp_path):
    md = tmp_path / "bom.md"
    md.write_bytes("# 带 BOM 的标题\n\n## 一\n\n- a\n- b\n".encode("utf-8-sig"))
    d, _ = profiles["finance-ruixin"]
    r = run(os.path.join(ENGINE, "plan.py"), str(md), str(d / "profile.json"), "--out", str(tmp_path / "deck.json"))
    assert r.returncode == 0, r.stderr
    assert json.load(open(tmp_path / "deck.json", encoding="utf-8"))["slides"][0]["title"] == "带 BOM 的标题"


def test_no_h2_keeps_content():
    doc = blocks("# T\n\n### 页一\n\n- **a**：1\n- **b**：2\n")
    assert doc["sections"] and doc["sections"][0]["units"], "没有 ## 的内容不能被丢掉"


def test_list_does_not_swallow_heading():
    b = parse("- a\n- b\n### Q\ntext\n")
    assert b[0]["items"] == ["a", "b"] and b[1]["t"] == "h"


def test_loose_list_is_one_block():
    b = parse("- a\n\n- b\n\n- c\n")
    assert len(b) == 1 and b[0]["items"] == ["a", "b", "c"]


def test_image_with_title_and_url():
    b = parse('![x](a.png "标题")\n\n![y](https://e.com/b.png)\n')
    assert [x.get("path") for x in b] == ["a.png"]


def test_long_table_paged():
    rows = "\n".join(f"| 项{i} | 说明{i} |" for i in range(20))
    doc = blocks(f"# T\n\n## S\n\n| 名称 | 说明 |\n|---|---|\n{rows}\n")
    out = unit_slides(doc["sections"][0]["units"][0])
    assert len(out) == 3 and all(o[0]["type"] == "table" and len(o[0]["rows"]) <= 8 for o in out)


def test_compare_followed_by_table_keeps_table():
    md = "# T\n\n## S\n\n**过去**\n\n- a\n- b\n\n**现在**\n\n- c\n- d\n\n| 年 | 值 |\n|---|---|\n| 2024 | 1 |\n| 2025 | 2 |\n"
    out = unit_slides(blocks(md)["sections"][0]["units"][0])
    assert [o[0]["type"] for o in out] == ["compare", "chart"]


def test_stats_uses_bold_number():
    out = unit_slides(blocks("# T\n\n## S\n\n- 国家 4A 级以上景区 **9 处**\n- 森林覆盖率 **78%**\n")["sections"][0]["units"][0])
    items = out[0][0]["items"]
    assert out[0][0]["type"] == "stats" and items[0]["value"] == "9" and items[0]["unit"] == "处"


@pytest.mark.parametrize("ex", ["finance", "medical", "education", "manufacturing", "tourism", "energy", "keynote",
                                "consulting"])
def test_examples_build_and_lint(ex, tmp_path):
    deck = os.path.join(ROOT, "examples", ex, "deck.json")
    out = tmp_path / "deck.pptx"
    r = run(os.path.join(ENGINE, "build.py"), deck, "--out", str(out), cwd=str(tmp_path))  # 在别的目录运行
    assert r.returncode == 0, r.stderr
    r = run(os.path.join(ENGINE, "lint.py"), str(out), "--no-render")
    assert r.returncode == 0, r.stdout


# ---------- 非 md 输入：convert.py ----------
def convert_to(src, tmp_path):
    from convert import convert
    out = tmp_path / "work" / "c.md"
    convert(str(src), str(out))
    return out.read_text(encoding="utf-8")


def test_txt_gbk_numbering(tmp_path):
    src = tmp_path / "a.txt"
    src.write_bytes("云栖推介\n\n2027 年 3 月\n\n第一章 资源\n景区 9 处\n覆盖率 78%\n\n1.1 四季线路\n春：赏花\n夏：避暑\n\n"
                    "第二章 合作\n• 包团\n• 研学\n项目\t2025\t2026\n游客\t320\t410\n".encode("gb18030"))
    md = convert_to(src, tmp_path)
    for line in ["# 云栖推介", "## 第一章 资源", "- 景区 9 处", "### 1.1 四季线路", "- 春：赏花", "- 包团",
                 "| 游客 | 320 | 410 |"]:
        assert line + "\n" in md, md
    assert "- 2027" not in md, "封面下面的日期不能被当成列表"


def test_txt_number_is_not_heading(tmp_path):
    src = tmp_path / "a.txt"
    src.write_text("标题\n\n一、背景\n2.5 亿元收入\n1.5 倍增长\n", encoding="utf-8")
    assert "### 2.5" not in convert_to(src, tmp_path)


def make_docx(path, styled):
    import docx
    from PIL import Image
    d = docx.Document()
    if styled:
        d.add_paragraph("康宁年度汇报", style="Title")
        d.add_heading("项目背景", 1)
        d.add_heading("为什么做", 2)
        d.add_paragraph("患者多：3.2 万人", style="List Bullet")
        d.add_paragraph("随访少：0.6 次", style="List Bullet")
        d.add_paragraph().add_run("结论：必须数字化。").bold = True
        d.add_heading("流程", 2)
        for t in ("建档", "评估", "干预"):
            d.add_paragraph(t, style="List Number")
        tb = d.add_table(rows=2, cols=2)
        for i, row in enumerate([["季度", "控制率"], ["Q1", "52%"]]):
            for j, v in enumerate(row):
                tb.cell(i, j).text = v
        img = path.parent / "p.png"
        Image.new("RGB", (40, 20), (0, 90, 160)).save(img)
        d.add_picture(str(img))
    else:
        d.add_paragraph().add_run("华铸工厂汇报").bold = True
        d.add_paragraph("一、建设背景")
        d.add_paragraph().add_run("产能瓶颈").bold = True
        d.add_paragraph("去年订单增长 40%，但老产线 OEE 只有 58%，交付经常延期。")
        d.add_paragraph("（一）设备联网")
        for t in ("1. 加装传感器", "2. 打通 MES"):
            d.add_paragraph(t)
    d.save(path)


def test_docx_heading_styles(tmp_path):
    src = tmp_path / "a.docx"
    make_docx(src, styled=True)
    md = convert_to(src, tmp_path)
    for line in ["# 康宁年度汇报", "## 项目背景", "### 为什么做", "- 患者多：3.2 万人", "**结论：必须数字化。**",
                 "1. 建档", "| Q1 | 52% |", "![](c-media/img01.png)"]:
        assert line + "\n" in md, md
    assert (tmp_path / "work" / "c-media" / "img01.png").exists()


def test_docx_without_styles(tmp_path):
    src = tmp_path / "a.docx"
    make_docx(src, styled=False)
    md = convert_to(src, tmp_path)
    for line in ["# 华铸工厂汇报", "## 一、建设背景", "### 产能瓶颈", "### （一）设备联网", "1. 加装传感器"]:
        assert line + "\n" in md, md


def test_pdf_lone_bullet_glyph():
    from convert import text_blocks
    b = text_blocks("要点\n\n\n第一条内容\n\n第二条内容\n")
    assert [x["text"] for x in b[1:]] == ["• 第一条内容", "• 第二条内容"]


def test_plan_accepts_docx(profiles, tmp_path):
    src = tmp_path / "汇报.docx"
    make_docx(src, styled=True)
    d, _ = profiles["medical-kangning"]
    r = run(os.path.join(ENGINE, "plan.py"), str(src), str(d / "profile.json"), "--out", str(tmp_path / "deck.json"))
    assert r.returncode == 0, r.stderr
    deck = json.load(open(tmp_path / "deck.json", encoding="utf-8"))
    assert deck["slides"][0]["title"] == "康宁年度汇报"
    assert (d / "汇报.converted.md").exists()
    assert any(s.get("body", {}).get("type") == "image" and os.path.exists(tmp_path / s["body"]["path"])
               for s in deck["slides"]), "docx 里的图片要能找到"


def test_sample_text_anywhere():
    from build import SAMPLE_RE
    assert SAMPLE_RE.search("SECTION SUBTITLE") and not SAMPLE_RE.search("华铸智造")


def txt_md(s, pdf=False):
    from convert import kv_lists, resolve_headings, text_blocks, to_md
    return to_md(kv_lists(resolve_headings(text_blocks(s, hard_wrapped=pdf))))


def test_txt_review_cases():
    assert "1. 加装传感器\n" in txt_md("标题\n\n一、背景\n1.\t加装传感器\n2.\t打通MES\n"), "编号+Tab 是列表不是表格"
    assert "|  | 2025 | 2026 |" in txt_md("标题\n\n一、背景\n\t2025\t2026\n游客\t320\t410\n"), "左上角空格子不能丢"
    assert "###" not in txt_md("标题\n\n一、背景\n\n4.8 分\n\n3.5个百分点\n"), "带单位的小数不是小节号"
    md = txt_md("标题\n\n一、背景\n游客人数\n3200\n2025\n- 3 -\n", pdf=True)
    assert "3200" in md and "2025" in md and "- 3 -" not in md, "PDF 只删页码，不删数据"
    assert "from the hunter" in txt_md("T\n\nPart 1\n" + "x" * 40 + " runs away from\nthe hunter ends here.\n", pdf=True)


def test_docx_review_cases(tmp_path):
    import docx
    from docx.oxml.ns import qn
    d = docx.Document()
    d.add_paragraph().add_run("华铸工厂汇报").bold = True  # 封面标题没用标题样式
    d.add_heading("建设背景", 1)
    p = d.add_paragraph("普通说明文字。", style="List Bullet")
    p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().set(qn("w:val"), "0")  # 取消编号
    p = d.add_paragraph("We saw the")
    p.add_run(" quick ").bold = True
    p.add_run("fox.")
    d.add_heading("实施", 1)
    tb = d.add_table(rows=3, cols=2)
    tb.cell(0, 0).merge(tb.cell(2, 0)).text = "合并"
    tb.cell(0, 1).text = "a"
    src = tmp_path / "a.docx"
    d.save(src)
    md = convert_to(src, tmp_path)
    for line in ["# 华铸工厂汇报", "## 建设背景", "普通说明文字。", "We saw the **quick** fox.", "|  | "]:
        assert line in md, md
    assert "1. 普通" not in md and "- 普通" not in md and md.count("合并") == 1


def test_converted_media_dir_is_md_safe(tmp_path):
    from convert import convert
    src = tmp_path / "a.docx"
    make_docx(src, styled=True)
    out = tmp_path / "年度 汇报(终稿).md"
    convert(str(src), str(out))
    assert "![](年度_汇报_终稿_-media/img01.png)" in out.read_text(encoding="utf-8")


def test_unfilled_section_subtitle_removed(tmp_path):
    out = tmp_path / "deck.pptx"
    r = run(os.path.join(ENGINE, "build.py"), os.path.join(ROOT, "examples", "education", "deck.json"), "--out", str(out))
    assert r.returncode == 0, r.stderr
    from pptx import Presentation
    texts = [sh.text_frame.text for s in Presentation(str(out)).slides for sh in s.shapes if sh.has_text_frame]
    assert "这一章你会学到什么" not in texts, "章节页没填副标题时，模板的示例副标题要删掉"


def test_muted_text_readable_on_cream():
    from common import contrast
    from draw import Theme
    tokens = {"primary": "2E7D6B", "secondary": "C0392B", "font": "x", "dark": "222222", "body": "444444",
              "muted": "8A8A8A", "radius": 0.1}
    assert contrast(Theme(tokens, "F0ECE1").muted, "F0ECE1") >= 3.0


# ---------- 复杂模板 ----------
@pytest.mark.parametrize("name", list(COMPLEX_ROLES))
def test_complex_roles(complex_profiles, name):
    """隐藏的说明页、议程式章节页、模板自带的金句页、左标题内容页都要认对"""
    assert [s["role"] for s in complex_profiles[name]["slides"]] == COMPLEX_ROLES[name]


def test_theme_brand_color(complex_profiles):
    """主色来自主题引用的品牌蓝，不是封面渐变遮罩的深色"""
    assert complex_profiles["energy-yuechuan"]["tokens"]["primary"] == "0B5CAD"


def test_tracker_and_footer(complex_profiles):
    nav_page = complex_profiles["energy-yuechuan"]["slides"][4]
    assert nav_page["tracker"]["axis"] == "col" and nav_page["tracker"]["current"] == 0 and nav_page["tracker"]["mark"]
    for s in complex_profiles["consulting-hengce"]["slides"][4:13]:
        assert s["tracker"]["axis"] == "row" and len(s["tracker"]["items"]) == 4
        assert "number" not in s["slots"], "页码占位符不能当成序号槽位"
    agenda = complex_profiles["consulting-hengce"]["slides"][3]
    assert agenda.get("agenda") and agenda["toc_mode"] == "paragraphs" and agenda["toc_current"] == 1


def test_paragraph_numbering():
    from nav import _numbered
    assert _numbered("1　　议题一：xx", 2) == "3　　" and _numbered("01. 标题", 3) == "04. " and \
        _numbered("一、背景", 1) == "二、"


def test_complex_build_fills_navigation(tmp_path):
    from pptx import Presentation
    out = tmp_path / "deck.pptx"
    r = run(os.path.join(ENGINE, "build.py"), os.path.join(ROOT, "examples", "consulting", "deck.json"), "--out",
            str(out))
    assert r.returncode == 0, r.stderr
    prs = Presentation(str(out))
    deck = json.load(open(os.path.join(ROOT, "examples", "consulting", "deck.json"), encoding="utf-8"))
    names = deck["slides"][1]["items"]
    for i, spec in enumerate(deck["slides"]):
        slide = prs.slides[i]
        if "current" in spec:  # 议程式章节页：段落是章节名，加粗的是当前章
            box = next(sh for sh in slide.shapes if sh.has_text_frame and len(sh.text_frame.paragraphs) >= 3)
            ps = box.text_frame.paragraphs
            assert [p.text.split("　")[-1] for p in ps] == names
            assert [bool(p.runs[0].font.bold) for p in ps].index(True) == spec["current"]
        if spec.get("tracker"):
            texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text in names
                     and not sh.is_placeholder and sh.top < 914400 * 0.6]  # 右上角的进度标签
            assert texts == names
    cover = prs.slides[0].shapes.title
    assert cover.top / 914400 > 1.5, "位置继承自版式的标题占位符，调整后不能跑到页面顶部"


def test_compare_fits_short_area(tmp_path):
    """版面很矮（10×5.625 英寸、带导航栏）的对比页：条目折行后也不能超出内容区"""
    prof = tmp_path / "work"
    r = run(os.path.join(ENGINE, "inspect_template.py"), "examples/templates/energy-yuechuan.pptx", "--out", str(prof),
            "--no-render")
    assert r.returncode == 0, r.stderr
    deck = {"template": os.path.join(ROOT, "examples/templates/energy-yuechuan.pptx"),
            "profile": str(prof / "profile.json"), "output": str(tmp_path / "out.pptx"),
            "slides": [{"from": 4, "title": "对比", "tracker": None, "area": [1.35, 0.91, 7.4, 2.9],
                        "body": {"type": "compare", "left": {"title": "过去", "items": ["靠人工巡检", "纸质工作票",
                                                                                         "事后追责"]},
                                 "right": {"title": "现在", "items": ["无人机和机器人巡检", "电子工作票全流程留痕",
                                                                      "事前风险预警"]}}}]}
    p = tmp_path / "deck.json"
    p.write_text(json.dumps(deck, ensure_ascii=False), encoding="utf-8")
    assert run(os.path.join(ENGINE, "build.py"), str(p)).returncode == 0
    r = run(os.path.join(ENGINE, "lint.py"), str(tmp_path / "out.pptx"), "--no-render")
    assert r.returncode == 0 and "超出内容区" not in r.stdout, r.stdout


def test_review_text_rules():
    from build import split_balanced
    from plan import NUM
    assert NUM.search("5千克").groups() == ("5", "千克") and NUM.search("300km").groups() == ("300", "km")
    assert NUM.search("2,860 亿千瓦时").group(2) == "亿千瓦时"
    assert split_balanced("数字化转型背景下，制造业供应链韧性提升") == ["数字化转型背景下，", "制造业供应链韧性提升"]
    assert split_balanced("星曜 X1 智能眼镜发布会") == ["星曜 X1", "智能眼镜发布会"]
    assert split_balanced("没有任何断点的一长串中文标题") == ["没有任何断点的一长串中文标题"]
    titles = [u["title"] for s in blocks("# T\n\n## S\n\n### 2.5 倍增长目标\n\n- a\n\n### 1.1 背景\n\n- b\n"
                                          "\n### （一）现状\n\n- c\n")["sections"] for u in s["units"]]
    assert titles == ["2.5 倍增长目标", "背景", "现状"]


def test_review_slot_rules():
    from inspect_template import assign_slots, classify

    def t(key, text, box, size, ph=None):
        return {"key": key, "text": text, "box": box, "size": size, "placeholder": ph, "wrap": True}
    texts = [t("#1", "章节标题", [1, 3, 6, 1], 40), t("#2", "SECTION SUBTITLE", [1, 4, 6, 0.4], 14),
             t("#3", "2026.12.18", [1, 5, 3, 0.4], 12), t("#4", "CHAPTER ONE", [1, 2, 3, 0.3], 10)]
    slots = assign_slots("section", texts, 13.33, 7.5)
    assert slots["kicker"] == ["#4"] and "#2" in slots["subtitle"] and slots["meta"] == ["#3"]
    page = [t("#2", "页面标题", [0.5, 0.4, 8, 0.6], 24, "TITLE"), t("#3", "", [0.5, 1.5, 12, 5], 32, "OBJECT")]
    assert classify(1, 0, 5, page, 13.33, 7.5) == "content", "空的正文占位符不能让普通页变成金句页"


def test_footer_row_is_not_tracker():
    from types import SimpleNamespace
    from lxml import etree
    from nav import detect_tracker
    A = "http://schemas.openxmlformats.org/drawingml/2006/main"

    def shape(bold):
        el = etree.fromstring(f'<sp xmlns:a="{A}"><a:r><a:rPr b="{1 if bold else 0}"/><a:t>x</a:t></a:r></sp>')
        return SimpleNamespace(_element=el, shape_type=17)

    def row(bolds, y):
        return [{"key": f"#{i}", "text": f"标签{i}", "box": [8 + i, y, 0.8, 0.3], "size": 9, "placeholder": None,
                 "shape": shape(b)} for i, b in enumerate(bolds)]
    assert detect_tracker(row([0, 0, 0], 7.1), [], 13.33, 7.5) is None, "没有高亮项的一排页脚小字不是进度标签"
    assert detect_tracker(row([0, 1, 0, 0], 0.2), [], 13.33, 7.5)["current"] == 1


def test_quote_attribution_split(complex_profiles, tmp_path):
    d = tmp_path / "w"
    d.mkdir()
    json.dump(complex_profiles["keynote-xingyao"], open(d / "profile.json", "w", encoding="utf-8"), ensure_ascii=False)
    md = tmp_path / "q.md"
    md.write_text("# T\n\n## A\n\n> 增长——不是目标，而是结果\n\n## B\n\n> 我们想做的，是一副眼镜。—— 创始人 林知远\n", encoding="utf-8")
    import shutil
    shutil.copy(os.path.join(ROOT, "examples/templates/keynote-xingyao.pptx"), d / "keynote-xingyao.pptx")
    prof = json.load(open(d / "profile.json", encoding="utf-8"))
    prof["template"] = "keynote-xingyao.pptx"
    json.dump(prof, open(d / "profile.json", "w", encoding="utf-8"), ensure_ascii=False)
    assert run(os.path.join(ENGINE, "plan.py"), str(md), str(d / "profile.json"), "--out",
               str(tmp_path / "deck.json")).returncode == 0
    quotes = [s for s in json.load(open(tmp_path / "deck.json", encoding="utf-8"))["slides"] if "金句" in s["_why"]]
    assert [(q["title"], q.get("subtitle")) for q in quotes] == [
        ("增长——不是目标，而是结果", None), ("我们想做的，是一副眼镜。", "—— 创始人 林知远")]
