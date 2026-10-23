"""公共工具：单位换算、颜色、文字宽度估算、形状遍历。所有脚本共用，不依赖字体文件，结果可复现。"""
import colorsys
import json
import os
import re
import unicodedata

from pptx.oxml.ns import qn

EMU = 914400  # 1 英寸
ENGINE = os.path.dirname(os.path.abspath(__file__))
ICONS = os.path.join(ENGINE, "icons")


def emu(inches):
    return int(round(inches * EMU))


def inch(e):
    return (e or 0) / EMU


def load_json(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def save_json(path, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


# ---------- 颜色 ----------
def norm_hex(c):
    c = c.strip().lstrip("#").upper()
    if len(c) == 3:
        c = "".join(ch * 2 for ch in c)
    if not re.fullmatch(r"[0-9A-F]{6}", c):
        raise ValueError(f"不是合法颜色：{c}")
    return c


def rgb(c):
    c = norm_hex(c)
    return tuple(int(c[i:i + 2], 16) for i in (0, 2, 4))


def to_hex(t):
    return "".join(f"{max(0, min(255, round(v))):02X}" for v in t)


def mix(c1, c2, t):
    """t=0 得到 c1，t=1 得到 c2"""
    a, b = rgb(c1), rgb(c2)
    return to_hex(a[i] + (b[i] - a[i]) * t for i in range(3))


def luminance(c):
    def ch(v):
        v /= 255
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = rgb(c)
    return 0.2126 * ch(r) + 0.7152 * ch(g) + 0.0722 * ch(b)


def contrast(c1, c2):
    a, b = sorted((luminance(c1), luminance(c2)), reverse=True)
    return (a + 0.05) / (b + 0.05)


def hls(c):
    r, g, b = (v / 255 for v in rgb(c))
    return colorsys.rgb_to_hls(r, g, b)  # (h, l, s)


def is_saturated(c):
    h, l, s = hls(c)
    return s > 0.35 and 0.18 < l < 0.82


def hue_dist(c1, c2):
    d = abs(hls(c1)[0] - hls(c2)[0])
    return min(d, 1 - d) * 360


def readable_on(bg, light="FFFFFF", dark="1F2329"):
    """在 bg 上取对比度更高的文字色"""
    return light if contrast(light, bg) >= contrast(dark, bg) else dark


# ---------- 文字宽度估算（单位 em，1em = 字号） ----------
def char_em(ch):
    if ch in "\t":
        return 1.0
    if unicodedata.east_asian_width(ch) in ("W", "F"):
        return 1.0
    if ch == " ":
        return 0.3
    if ch in "il.,:;'!|`()[]{}·":
        return 0.32
    if ch in "mwMW@%":
        return 0.85
    if ch.isupper():
        return 0.7
    if ch.isdigit():
        return 0.6
    return 0.52


def text_em(s):
    return sum(char_em(ch) for ch in s)


def text_width(s, size_pt):
    """一行文字的估算宽度（英寸）"""
    return text_em(s) * size_pt / 72


WRAP_FUDGE = 1.08  # 实际字体（尤其西文、标点挤压）比估算宽一些，按 8% 余量折行


def _tokens(line):
    """中文逐字可断行，英文/数字按单词断行"""
    return re.findall(r"[A-Za-z0-9][A-Za-z0-9.,%+\-/'’]*\s*|\s+|.", line)


def count_lines(text, width_in, size_pt):
    if width_in <= 0:
        return 999
    cap = width_in * 72 / size_pt / WRAP_FUDGE  # 一行能放多少 em
    total = 0
    for para in str(text).split("\n"):
        lines, cur = 1, 0.0
        for tok in _tokens(para):
            w = text_em(tok)
            if cur + w > cap + 1e-6 and cur > 0:
                lines += 1
                cur = text_em(tok.lstrip())
                while cur > cap:  # 超长单词硬断
                    lines += 1
                    cur -= cap
            else:
                cur += w
        total += lines
    return total


def text_height(text, width_in, size_pt, spacing=1.2):
    return count_lines(text, width_in, size_pt) * size_pt * spacing / 72


def fit_size(text, w, h, size, min_size, spacing=1.2):
    """在 w×h 里放得下的最大字号（每次减 1pt），放不下返回 None"""
    s = size
    while s >= min_size:
        if text_height(text, w, s, spacing) <= h + 0.02:
            return s
        s -= 1
    return None


# ---------- 形状遍历与几何 ----------
def iter_shapes(shapes, xf=None):
    """遍历所有形状（含组合内部），返回 (shape, (x, y, w, h) 英寸绝对坐标, 是否在组合里)"""
    for s in shapes:
        box = abs_box(s, xf)
        yield s, box, xf is not None
        if s.shape_type == 6:  # GROUP
            yield from iter_shapes(s.shapes, group_xf(s, xf))


def group_xf(g, parent):
    x = g._element.grpSpPr.find(qn("a:xfrm"))
    off, ext = x.find(qn("a:off")), x.find(qn("a:ext"))
    choff, chext = x.find(qn("a:chOff")), x.find(qn("a:chExt"))
    sx = int(ext.get("cx")) / max(1, int(chext.get("cx")))
    sy = int(ext.get("cy")) / max(1, int(chext.get("cy")))
    t = (int(off.get("x")), int(off.get("y")), int(choff.get("x")), int(choff.get("y")), sx, sy)
    return (parent, t)


def abs_box(s, xf):
    try:
        x, y, w, h = s.left or 0, s.top or 0, s.width or 0, s.height or 0
    except Exception:
        x = y = w = h = 0
    while xf is not None:
        parent, (ox, oy, cx, cy, sx, sy) = xf
        x, y, w, h = ox + (x - cx) * sx, oy + (y - cy) * sy, w * sx, h * sy
        xf = parent
    return (inch(x), inch(y), inch(w), inch(h))


def overlap(a, b):
    """两个矩形 (x,y,w,h) 的交叠面积"""
    dx = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
    dy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
    return dx * dy if dx > 0 and dy > 0 else 0.0


def shape_text(s):
    if not getattr(s, "has_text_frame", False) or not s.has_text_frame:
        return ""
    return "\n".join(p.text for p in s.text_frame.paragraphs).strip()


def first_size(s, default=18.0):
    """形状里第一个显式字号（pt）"""
    for tag in ("a:rPr", "a:defRPr", "a:endParaRPr"):
        for el in s._element.iter(qn(tag)):
            if el.get("sz"):
                return int(el.get("sz")) / 100
    return default


def max_size(s, default=18.0):
    sizes = [int(el.get("sz")) / 100 for el in s._element.iter(qn("a:rPr")) if el.get("sz")]
    return max(sizes) if sizes else first_size(s, default)


def wrap_none(s):
    bp = s._element.find(".//" + qn("a:bodyPr"))
    return bp is not None and bp.get("wrap") == "none"


def text_insets(s):
    """文本框内边距（英寸）：左、右、上、下"""
    bp = s._element.find(".//" + qn("a:bodyPr"))
    d = {"lIns": 91440, "rIns": 91440, "tIns": 45720, "bIns": 45720}
    if bp is not None:
        for k in d:
            if bp.get(k) is not None:
                d[k] = int(bp.get(k))
    return tuple(inch(d[k]) for k in ("lIns", "rIns", "tIns", "bIns"))


def slot_key(s):
    """在 deck.json 里引用形状的方式：#形状 id"""
    return f"#{s.shape_id}"


def profile_template(prof, prof_path):
    """profile.json 里的模板路径是相对 profile 所在目录存的（换机器也能用）"""
    t = prof["template"]
    return t if os.path.isabs(t) else os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(prof_path)), t))
