"""不用看图的成品体检。用法：
    python3 lint.py out/deck.pptx [--no-render]

检查项（❌ 必须修，⚠️ 建议修）：
  ❌ 文字溢出（按字号估算行数，超出文本框）      ❌ 形状出界（新画的内容超出页面）
  ❌ 文字互相压住（新文字之间、新文字压模板文字）  ❌ 模板示例字残留（XXXX / 标题 / 请输入…，或标题槽没改）
  ❌ 字号小于 10pt                               ❌ 文字和背景对比度 < 3（需要渲染，按真实像素算）
  ❌ 内容页空着 / 新内容超出内容区下沿          ⚠️ 单页字数过多 / 使用了模板以外的字体 / 构建时的提示
  （对比度：模板自带的文字原样保留了模板颜色时只给 ⚠️，不要为了过检改模板的颜色）
有 ❌ 时退出码为 1。
"""
import argparse
import collections
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

from build import SAMPLE_RE
from inspect_template import ph_size
from common import profile_template, contrast, count_lines, inch, iter_shapes, load_json, overlap, save_json, text_em, text_insets, \
    to_hex, wrap_none


def runs_info(shape):
    sizes, colors, fonts = [], [], set()
    for r in shape._element.iter(qn("a:rPr")):
        if r.get("sz"):
            sizes.append(int(r.get("sz")) / 100)
        c = r.find(qn("a:solidFill") + "/" + qn("a:srgbClr"))
        if c is not None:
            colors.append(c.get("val").upper())
        for t in (qn("a:ea"), qn("a:latin")):
            f = r.find(t)
            if f is not None and f.get("typeface") and not f.get("typeface").startswith("+"):
                fonts.add(f.get("typeface"))
    return sizes, colors, fonts


def para_width(p, default):
    """按每个 run 自己的字号算一段的宽度（英寸）：大数字 + 小单位不会被高估"""
    w = 0.0
    for r in p.runs:
        sz = r.font.size.pt if r.font.size else default
        w += text_em(r.text) * sz / 72
    return w


def overflow(shape, box, size):
    paras = shape.text_frame.paragraphs
    lines = [p.text for p in paras]
    l_in, r_in, t_in, b_in = text_insets(shape)
    widths = [para_width(p, size) for p in paras]
    if wrap_none(shape):
        need = max(widths) + l_in + r_in
        return need > box[2] * 1.12 + 0.1, f"一行需要 {need:.1f} 英寸，框宽 {box[2]:.1f}"
    if len(paras) == 1 and len(paras[0].runs) > 1:  # 混合字号的单段：按宽度折行
        nl = max(1, -(-widths[0] * 1.08 // max(0.1, box[2] - l_in - r_in)))
        need = nl * size * 1.2 / 72 + t_in + b_in
        return need > box[3] * 1.15 + 0.12, f"需要约 {int(nl)} 行 / {need:.1f} 英寸高，框高 {box[3]:.1f}"
    autofit = shape._element.find(".//" + qn("a:spAutoFit")) is not None
    nl = sum(count_lines(t, box[2] - l_in - r_in, size) for t in lines)
    need = nl * size * 1.2 / 72 + t_in + b_in
    lim = box[3] * (1.6 if autofit else 1.15) + 0.12
    return need > lim, f"需要约 {nl} 行 / {need:.1f} 英寸高，框高 {box[3]:.1f}"


def bg_contrast(img, W, H, box, color):
    sx, sy = img.width / W, img.height / H
    x, y, w, h = box
    crop = img.crop((int(x * sx), int(y * sy), int((x + w) * sx) + 1, int((y + h) * sy) + 1))
    if crop.width < 2 or crop.height < 2:
        return None
    crop = crop.resize((min(60, crop.width), min(30, crop.height)))
    cnt = collections.Counter(crop.getdata())
    tc = tuple(int(color[i:i + 2], 16) for i in (0, 2, 4))
    # 背景 = 出现最多、且和文字色差得最远的那批像素
    cands = [c for c, _ in cnt.most_common(6) if sum(abs(c[i] - tc[i]) for i in range(3)) > 60]
    if not cands:
        return None
    return contrast(color, to_hex(cands[0])), to_hex(cands[0])


def lint(path, do_render=True):
    rep_path = path + ".build.json"
    rep = load_json(rep_path) if os.path.exists(rep_path) else None
    if rep and not os.path.isabs(rep["profile"]):
        rep["profile"] = os.path.join(os.path.dirname(os.path.abspath(path)), rep["profile"])
    prof = load_json(rep["profile"]) if rep and os.path.exists(rep["profile"]) else None
    tpl = Presentation(profile_template(prof, rep["profile"])) if prof else None
    prs = Presentation(path)
    W, H = inch(prs.slide_width), inch(prs.slide_height)
    tpl_fonts = set()
    if tpl is not None:
        tpl_fonts = {prof["tokens"]["font"], prof["tokens"]["font_heading"], prof["tokens"]["font_latin"]}
        for sl in tpl.slides:
            for s, _, _ in iter_shapes(sl.shapes):
                if s.has_text_frame:
                    tpl_fonts |= runs_info(s)[2]
    imgs = []
    if do_render:
        from render import render
        out = os.path.join(os.path.dirname(os.path.abspath(path)), "lint-render")
        imgs = [Image.open(f).convert("RGB") for f in render(path, out, dpi=50)]
    issues = []

    def add(page, level, msg):
        issues.append({"page": page, "level": level, "msg": msg})

    for i, slide in enumerate(prs.slides):
        page = i + 1
        info = rep["slides"][i] if rep and i < len(rep["slides"]) else {}
        fi = info.get("from")
        tpl_ids, tpl_colors = set(), {}
        if tpl is not None and fi is not None:
            tpl_ids = {s.shape_id for s, _, _ in iter_shapes(tpl.slides[fi].shapes)}
            tpl_colors = {s.shape_id: runs_info(s)[1] for s, _, _ in iter_shapes(tpl.slides[fi].shapes) if s.has_text_frame}
        ps = prof["slides"][fi] if prof and fi is not None else None
        tpl_texts = {t["text"] for t in ps["texts"]} if ps else set()
        new_text, old_text, chars = [], [], 0
        new_count = 0
        spill = 0.0  # 新内容超出内容区下沿多少
        for s, box, nested in iter_shapes(slide.shapes):
            is_new = s.name.startswith("xm ") if rep else s.shape_id not in tpl_ids
            new_count += is_new
            if is_new and s.shape_type != 6 and (box[0] < -0.02 or box[1] < -0.02 or box[0] + box[2] > W + 0.02
                                                 or box[1] + box[3] > H + 0.02):
                add(page, "error", f"形状出界：{s.name} 位置 {tuple(round(v, 2) for v in box)}")
            ar = info.get("area")
            if is_new and ar and s.shape_type != 6 and box[1] + box[3] > ar[1] + ar[3] + 0.15:
                spill = max(spill, box[1] + box[3] - ar[1] - ar[3])
            if is_new and getattr(s, "has_table", False) and s.has_table:
                rows = len(s.table.rows)
                min_h = rows * 0.3
                if box[1] + max(box[3], min_h) > H + 0.05:
                    add(page, "error", f"表格 {rows} 行放不下（至少需要 {min_h:.1f} 英寸高），请分页")
            if not s.has_text_frame or not s.text_frame.text.strip():
                continue
            txt = s.text_frame.text.strip()
            chars += len(txt)
            sizes, colors, fonts = runs_info(s)
            size = max(sizes) if sizes else (ph_size(s, slide) if s.is_placeholder else 18)  # 占位符的字号继承版式 / 母版
            short = txt.replace("\n", " ")[:20]
            if is_new and sizes and min(sizes) < 10:
                add(page, "error", f"字号 {min(sizes)}pt 太小：「{short}」")
            bad, why = overflow(s, box, size) if box[2] > 0 and box[3] > 0 else (False, "")  # 位置继承母版的占位符
            if bad:
                add(page, "error", f"文字溢出：「{short}」{why}")
            if not is_new and txt in tpl_texts and SAMPLE_RE.search(txt) and not txt.strip().isdigit():
                add(page, "error", f"模板示例字没替换：「{short}」")
            for f in fonts - tpl_fonts if tpl_fonts else ():
                add(page, "warn", f"用了模板以外的字体 {f}：「{short}」")
            if imgs and colors and i < len(imgs):
                r = bg_contrast(imgs[i], W, H, box, colors[0])
                if r and r[0] < 3.0 and not is_new and colors[0] in tpl_colors.get(s.shape_id, ()):
                    # 模板自己就是这个配色（例如白底金色序号）：规矩 1 是不改模板颜色，只提醒
                    add(page, "warn", f"对比度 {r[0]:.1f} 偏低，但这是模板自带的配色（文字 #{colors[0]} / 背景 #{r[1]}），"
                                      f"保留不改：「{short}」")
                elif r and r[0] < 3.0:
                    add(page, "error", f"对比度 {r[0]:.1f}（文字 #{colors[0]} / 背景 #{r[1]}）看不清：「{short}」")
            (new_text if is_new else old_text).append((s, box, short))
        if spill:
            add(page, "error", f"新内容超出内容区下沿 {spill:.2f} 英寸（压到了模板的页脚或装饰）：删减内容或拆页")
        if ps and ps["role"] == "content" and new_count == 0:
            add(page, "error", "内容页没有任何内容（缺 body）")
        if ps and ps["slots"].get("title") and ps["role"] in ("cover", "section", "content"):
            ids = {f"#{s.shape_id}" for s, _, _ in iter_shapes(slide.shapes)}
            if ps["slots"]["title"] not in ids and info.get("body") not in ("quote", "statement"):
                add(page, "error", "标题槽空着（deck.json 里没写 title）")
        if ps and ps["slots"].get("title"):
            t0 = next((t["text"] for t in ps["texts"] if t["key"] == ps["slots"]["title"]), None)
            now = [s.text_frame.text.strip() for s, _, _ in iter_shapes(slide.shapes)
                   if f"#{s.shape_id}" == ps["slots"]["title"] and s.has_text_frame]
            if t0 and now and now[0] == t0 and ps["role"] != "ending":
                add(page, "error", f"标题还是模板原文「{t0[:16]}」")
        for a in range(len(new_text)):
            for b in range(a + 1, len(new_text)):
                sa, ba, ta = new_text[a]
                sb, bb, tb = new_text[b]
                ov = overlap(ba, bb)
                if ov > 0.15 * min(ba[2] * ba[3], bb[2] * bb[3]) and not contains(ba, bb) and not contains(bb, ba):
                    add(page, "error", f"文字压在一起：「{ta}」和「{tb}」")
            for so, bo, to in old_text:
                if overlap(new_text[a][1], bo) > 0.1 * min(bo[2] * bo[3], new_text[a][1][2] * new_text[a][1][3]):
                    add(page, "error", f"新内容压住了模板文字「{to}」：「{new_text[a][2]}」")
        if chars > 260:
            add(page, "warn", f"这一页有 {chars} 个字，太密：拆页或把细节挪到备注")
        for m in info.get("messages", []):
            add(page, "warn", f"构建提示：{m}")
    return issues


def contains(a, b):
    """a 完全包住 b（例如色块里的文字），不算压字"""
    return a[0] <= b[0] + 0.02 and a[1] <= b[1] + 0.02 and a[0] + a[2] >= b[0] + b[2] - 0.02 and \
        a[1] + a[3] >= b[1] + b[3] - 0.02


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx")
    ap.add_argument("--no-render", action="store_true", help="不渲染（跳过对比度检查）")
    a = ap.parse_args()
    issues = lint(a.pptx, not a.no_render)
    save_json(a.pptx + ".lint.json", issues)
    errs = [x for x in issues if x["level"] == "error"]
    warns = [x for x in issues if x["level"] == "warn"]
    for x in sorted(issues, key=lambda x: (x["page"], x["level"])):
        print(f"{'❌' if x['level'] == 'error' else '⚠️ '} 第 {x['page']} 页  {x['msg']}")
    print(f"\n{'全绿 ✅' if not errs else '有问题'}：{len(errs)} 个 ❌，{len(warns)} 个 ⚠️")
    sys.exit(1 if errs else 0)


if __name__ == "__main__":
    main()
