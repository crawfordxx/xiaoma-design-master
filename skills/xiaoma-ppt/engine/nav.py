"""导航类结构：真实企业模板里很常见，又不是普通的「标题 / 正文」槽位。
- 进度标签 / 章节导航条（tracker）：内容页边上一排或一列短标签，当前章节加粗、变色或垫底色
- 段落式目录：一个文本框里多个段落，每段一个议题
- 议程式章节页：高亮当前议题的目录页，当章节分隔页用
inspect_template.py 用 detect_*，build.py 用 fill_*。
"""
import copy
import re

from pptx.oxml.ns import qn

from common import emu, inch, overlap, text_em, text_insets

PREFIX = re.compile(r"^(\s*)(\d{1,2}|[一二三四五六七八九十]+)([\s.、．:：　\t]*)")
CN_NUM = "一二三四五六七八九十"


# ---------- 样式指纹 ----------
def rpr_sig(rpr):
    """一段文字的「样子」：加粗 + 颜色。用来找出被高亮的那一项"""
    if rpr is None:
        return (None, None)
    clr = rpr.find(qn("a:solidFill"))
    c = None
    if clr is not None and len(clr):
        c = clr[0].get("val")
    return (rpr.get("b"), c)


def first_rpr(el):
    r = el.find(".//" + qn("a:r"))
    return r.find(qn("a:rPr")) if r is not None else None


def highlight(sigs):
    """恰好一项和其余各项不同 → 返回它的序号"""
    if len(sigs) < 3:
        return None
    for i, s in enumerate(sigs):
        others = sigs[:i] + sigs[i + 1:]
        if s not in others and len(set(others)) == 1:
            return i
    return None


# ---------- 识别 ----------
def detect_tracker(texts, shapes, W, H):
    """内容页边缘一排（或一列）≥3 个等大短标签 → {items, current, mark, axis}"""
    cands = [t for t in texts if t["text"] and "\n" not in t["text"] and len(t["text"]) <= 12
             and not t["placeholder"]]
    best = None
    for axis in ("row", "col"):
        used = set()
        for t in cands:
            if t["key"] in used:
                continue
            c = _cross(t, axis)
            grp = [u for u in cands if abs(_cross(u, axis) - c) < 0.1 and abs(u["size"] - t["size"]) < 0.6]
            used |= {u["key"] for u in grp}
            if len(grp) < 3:
                continue
            grp.sort(key=lambda u: u["box"][0] if axis == "row" else u["box"][1])
            b0 = [u["box"] for u in grp]
            if axis == "row":
                edge = max(b[1] + b[3] for b in b0) < H * 0.2 or min(b[1] for b in b0) > H * 0.85
                small = all(b[2] < W * 0.2 for b in b0)
            else:
                edge = max(b[0] + b[2] for b in b0) < W * 0.16 or min(b[0] for b in b0) > W * 0.84
                small = all(b[3] < H * 0.2 for b in b0)
            # 必须恰好一格被高亮：没有高亮的一排小字多半是页脚（公司名 / 保密 / 年份）
            cur = highlight([rpr_sig(first_rpr(u["shape"]._element)) for u in grp])
            if edge and small and cur is not None and (best is None or len(grp) > len(best[1])):
                best = (axis, grp, cur)
    if not best:
        return None
    axis, grp, cur = best
    mark = _mark(grp[cur]["box"], shapes, W, H, axis, {u["key"] for u in grp})
    return {"items": [u["key"] for u in grp], "current": cur, "mark": mark, "axis": axis}


def _cross(t, axis):
    b = t["box"]
    return b[1] + b[3] / 2 if axis == "row" else b[0] + b[2] / 2


def _mark(cb, shapes, W, H, axis, skip):
    """当前项的标记：垫在它下面的色块，或紧贴它下方的短横线"""
    best = None
    for d in shapes:
        if d["key"] in skip or d["text"] or d["shape"].shape_type in (6, 13):
            continue
        b = d["box"]
        if b[2] * b[3] > W * H * 0.03 or b[2] * b[3] <= 0:
            continue
        behind = overlap(b, cb) >= 0.5 * cb[2] * cb[3]
        under = cb[1] + cb[3] - 0.08 <= b[1] <= cb[1] + cb[3] + 0.2 and overlap(
            (b[0], 0, b[2], 1), (cb[0], 0, cb[2], 1)) >= 0.5 * b[2]
        if (behind or under) and (best is None or b[2] * b[3] < best[1]):
            best = (d["key"], b[2] * b[3])
    return best[0] if best else None


def para_items(t):
    """一个文本框里 ≥3 个短段落 → 段落式目录"""
    ps = [p for p in t["shape"].text_frame.paragraphs if p.text.strip()]
    if len(ps) < 3 or any(len(p.text) > 40 for p in ps):
        return None
    cur = highlight([rpr_sig(first_rpr(p._p)) for p in ps])
    return {"count": len(ps), "current": cur}


def para_mark(t, shapes, current):
    """段落式议程里，标在当前段落旁边的竖条"""
    if current is None:
        return None
    b = t["box"]
    for d in shapes:
        db = d["box"]
        if d["key"] == t["key"] or d["text"] or d["shape"].shape_type in (6, 13):
            continue
        if db[2] < 0.3 and db[3] < 0.8 and b[0] - 0.6 <= db[0] <= b[0] + 0.1 and b[1] - 0.1 <= db[1] <= b[1] + b[3]:
            return d["key"]
    return None


# ---------- 填写 ----------
def _set_run(p_el, text, rpr):
    runs = p_el.findall(qn("a:r"))
    if not runs:
        r = p_el.makeelement(qn("a:r"), {})
        r.append(p_el.makeelement(qn("a:t"), {}))
        end = p_el.find(qn("a:endParaRPr"))
        (end.addprevious(r) if end is not None else p_el.append(r))
        runs = [r]
    for r in runs[1:]:
        p_el.remove(r)
    for ch in p_el.findall(qn("a:br")) + p_el.findall(qn("a:fld")):
        p_el.remove(ch)
    r = runs[0]
    if rpr is not None:
        old = r.find(qn("a:rPr"))
        if old is not None:
            r.remove(old)
        r.insert(0, copy.deepcopy(rpr))
    r.find(qn("a:t")).text = text


def _fit(shape, box, text):
    """标签格子很窄：字放不下就缩小字号（最小 6pt）"""
    rpr = first_rpr(shape._element)
    if rpr is None:
        return
    size = int(rpr.get("sz", "1200")) / 100
    li, ri, _, _ = text_insets(shape)
    room = box[2] - li - ri - 0.02
    need = text_em(text) * size / 72 * 1.1
    if need > room > 0:
        rpr.set("sz", str(int(max(6, size * room / need) * 100)))


def _move(slide, find_shape, key, d_in, axis):
    s, b, _ = find_shape(slide, key)
    if s is None or not d_in:
        return
    k = (b[2] / inch(s.width)) if axis == "row" and s.width else ((b[3] / inch(s.height)) if s.height else 1)
    if axis == "row":
        s.left = s.left + emu(d_in / max(k, 1e-6))
    else:
        s.top = s.top + emu(d_in / max(k, 1e-6))


def norm_current(current, n):
    """current 规整成 0..n-1 的整数，不合法就是 None（没有当前项）"""
    try:
        c = int(current)
    except (TypeError, ValueError):
        return None
    return c if 0 <= c < n else None


def _off_rpr(rpr):
    """没有可参照的非高亮样式时：去掉加粗和颜色"""
    r = copy.deepcopy(rpr)
    r.attrib.pop("b", None)
    for t in ("a:solidFill", "a:gradFill"):
        el = r.find(qn(t))
        if el is not None:
            r.remove(el)
    return r


def styles(els, i0):
    """模板里高亮项和普通项的文字样式 (on, off)；没有高亮项时都是 None"""
    if i0 is None or i0 >= len(els) or els[i0] is None:
        return None, None
    on = first_rpr(els[i0])
    off = next((first_rpr(e) for k, e in enumerate(els) if k != i0 and e is not None and first_rpr(e) is not None),
               None)
    if on is None:
        return None, None
    return copy.deepcopy(on), copy.deepcopy(off) if off is not None else _off_rpr(on)


def fill_tracker(slide, info, items, current, find_shape, remove, log, label="进度标签"):
    """把章节名写进进度标签；当前章节用模板高亮样式，标记跟着移过去。
    items 为空 → 整组删掉；current 为 None（附录、章节比格子多）→ 都用普通样式，标记删掉"""
    found = [find_shape(slide, k) for k in info["items"]]
    mark = find_shape(slide, info["mark"])[0] if info.get("mark") else None
    if not items:
        for s, _, _ in found:
            if s is not None:
                remove(s)
        if mark is not None:
            remove(mark)
        return
    if len(items) > len(found):
        log.append(f"{label}只有 {len(found)} 格，章节有 {len(items)} 个：只放了前 {len(found)} 个")
    n = min(len(items), len(found))
    cur = norm_current(current, n)
    if current is not None and cur is None:
        log.append(f"{label}的 current={current!r} 不在 0–{n - 1} 之间，按没有当前项处理")
    i0 = info.get("current")
    on, off = styles([s._element if s is not None else None for s, _, _ in found], i0)
    for j, (s, box, _) in enumerate(found):
        if s is None:
            continue
        if j >= len(items):
            remove(s)
            continue
        ps = s.text_frame._txBody.findall(qn("a:p"))
        for p in ps[1:]:
            p.getparent().remove(p)
        rpr = (on if j == cur else off) if on is not None else None
        _set_run(ps[0], items[j], rpr)
        _fit(s, box, items[j])
    if mark is not None:
        if cur is None or i0 is None or found[i0][0] is None or found[cur][0] is None:
            remove(mark)
        else:
            a, b = found[i0][1], found[cur][1]
            axis = info.get("axis", "row")
            _move(slide, find_shape, info["mark"], (b[0] - a[0]) if axis == "row" else (b[1] - a[1]), axis)


def _numbered(orig, j):
    m = PREFIX.match(orig)
    if not m:
        return ""
    n = m.group(2)
    if n.isdigit():
        num = f"{j + 1:0{len(n)}d}" if n.startswith("0") else str(j + 1)
    else:
        num = CN_NUM[j] if j < 10 else str(j + 1)
    return m.group(1) + num + m.group(3)


def _pitch(p_el):
    """一个段落占多高（英寸）：字号 × 行距 + 段前段后"""
    rpr = first_rpr(p_el)
    size = int(rpr.get("sz", "1800")) / 100 if rpr is not None and rpr.get("sz") else 18
    ppr = p_el.find(qn("a:pPr"))
    ls, after, before = 1.2, 0, 0
    if ppr is not None:
        pct = ppr.find(qn("a:lnSpc") + "/" + qn("a:spcPct"))
        if pct is not None:
            ls = int(pct.get("val")) / 100000 * 1.2
        for tag, attr in (("a:spcAft", "after"), ("a:spcBef", "before")):
            pts = ppr.find(qn(tag) + "/" + qn("a:spcPts"))
            if pts is not None:
                if attr == "after":
                    after = int(pts.get("val")) / 100
                else:
                    before = int(pts.get("val")) / 100
    return (size * ls + after + before) / 72


def _ptext(p):
    return "".join(t.text or "" for t in p.iter(qn("a:t")))


def fill_paragraphs(slide, key, items, current, i0, mark, find_shape, log):
    """段落式目录 / 议程：每个议题一段，保留模板的编号写法和段间空行；current 是要高亮的议题"""
    s, _, _ = find_shape(slide, key)
    if s is None:
        log.append(f"找不到目录文本框 {key}")
        return
    body = s.text_frame._txBody
    allp = body.findall(qn("a:p"))
    ps = [p for p in allp if _ptext(p).strip()]
    if not ps:
        return
    gaps = [p for p in allp if not _ptext(p).strip()]
    spacer = gaps[0] if gaps and allp.index(gaps[0]) > allp.index(ps[0]) else None  # 模板用空段落隔开议题
    texts = [_ptext(p) for p in ps]
    cur = norm_current(current, len(items))
    on, off = styles(ps, i0)
    if len(items) > len(ps):
        log.append(f"议程模板只有 {len(ps)} 段，章节有 {len(items)} 个：多出的按最后一段的样子补上")
    new = []
    for j, item in enumerate(items):
        p = copy.deepcopy(ps[min(j, len(ps) - 1)])
        rpr = (on if j == cur else off) if on is not None else first_rpr(p)
        _set_run(p, _numbered(texts[min(j, len(texts) - 1)], j) + item, rpr)
        if new and spacer is not None:
            new.append(copy.deepcopy(spacer))
        new.append(p)
    for p in allp:
        body.remove(p)
    for p in new or [copy.deepcopy(ps[0])]:  # 至少留一段，空 txBody 会让 PowerPoint 报修复
        body.append(p)
    if not new:
        _set_run(body.find(qn("a:p")), "", None)
    m = find_shape(slide, mark)[0] if mark else None
    if m is not None:
        if cur is None or i0 is None:
            m._element.getparent().remove(m._element)
        else:
            pitch = _pitch(ps[i0]) + (_pitch(spacer) if spacer is not None else 0)
            m.top = m.top + emu((cur - i0) * pitch)
