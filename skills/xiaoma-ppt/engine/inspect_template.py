"""拆解 PPT 模板，输出模板画像（profile）。用法：
    python3 inspect_template.py template.pptx --out work/            # 生成 profile.json / profile.md / template-sheet.jpg
    python3 inspect_template.py template.pptx --out work/ --no-render  # 没装 LibreOffice 时只做几何分析

画像包含：设计令牌（主色、辅色、文字色、字体、圆角）、每页角色（cover/toc/section/content/ending/other）、
每页可填的文字槽位（#形状id）、内容页的「内容安全区」、要删掉的示例内容。模型读 profile.md 就能写 deck.json。
"""
import argparse
import collections
import os
import re
import sys

from PIL import Image, ImageFilter
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn

import nav
from common import (first_size, hls, inch, is_saturated, iter_shapes, luminance, max_size, mix, norm_hex, save_json, shape_text, to_hex, wrap_none, hue_dist)

TOC_RE = re.compile(r"^\s*(目\s*录|contents?|agenda|议\s*程|catalog|目\s*录\s*/?\s*contents?)\s*$", re.I | re.M)
END_RE = re.compile(r"谢\s*谢|感\s*谢|THANKS?|Thanks?|Thank you|THANK YOU|再\s*见|Q\s*&\s*A|敬请|期待.*合作")
NUM_RE = re.compile(r"(PART|Part|CHAPTER|Chapter|SECTION|第)?\s*0?\d{1,2}\s*(章|部分|节)?\s*[.、]?")
DATE_RE = re.compile(r"(19|20)[\dXx]{2}\s*[年./\-]|\d{1,2}\s*月|汇报人|演讲人|主讲|日期|DATE", re.I)
TITLE_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
FOOTER_TYPES = (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE)
DISCLAIMER_RE = re.compile(r"免责|版权所有|未经.{0,12}许可|不得.{0,8}(披露|转载|复制)|disclaimer|copyright|©|all rights reserved",
                           re.I)
QUOTE_MARK = re.compile(r"\s*[“”\"「」『』❝❞]\s*")
CAPS = re.compile(r"[A-Z0-9][A-Z0-9 &·.,'’\-/|]{3,}")  # 全大写英文小字：装饰性的 kicker，不当副标题
ATTRIB = re.compile(r"^\s*(——|—|--|–)")


# ---------- 主题 ----------
def theme_info(prs):
    master = prs.slide_master
    theme_part = next((r.target_part for r in master.part.rels.values() if r.reltype.endswith("/theme")), None)
    fonts = {k: {"latin": "", "ea": ""} for k in ("majorFont", "minorFont")}
    if theme_part is None:
        return {}, fonts
    from lxml import etree
    root = etree.fromstring(theme_part.blob)
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    colors = {}
    scheme = root.find(f".//{a}clrScheme")
    for el in (scheme if scheme is not None else []):
        if not isinstance(el.tag, str) or len(el) == 0:  # 跳过注释
            continue
        name = el.tag.replace(a, "")
        c = el[0]
        try:
            colors[name] = norm_hex(c.get("val") if c.tag.endswith("srgbClr") else c.get("lastClr", "000000"))
        except (ValueError, AttributeError):
            pass
    for kind in ("majorFont", "minorFont"):
        f = root.find(f".//{a}{kind}")
        if f is None:
            continue
        for script in ("latin", "ea"):
            x = f.find(f"{a}{script}")
            fonts[kind][script] = (x.get("typeface") or "") if x is not None else ""
        for s in f.findall(f"{a}font"):
            if s.get("script") == "Hans" and not fonts[kind]["ea"]:
                fonts[kind]["ea"] = s.get("typeface")
    return colors, fonts


SCHEME_ALIAS = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}


def resolve_font(tf, fonts):
    if tf.startswith("+mj"):
        return fonts["majorFont"]["ea" if "ea" in tf else "latin"]
    if tf.startswith("+mn"):
        return fonts["minorFont"]["ea" if "ea" in tf else "latin"]
    return tf


# ---------- 设计令牌 ----------
def _skip(el):
    """p:style（形状默认样式引用主题色）和 p:txStyles（母版默认文字样式）不代表模板的真实用色"""
    p = el.getparent()
    while p is not None:
        if p.tag in (qn("p:style"), qn("p:txStyles"), qn("p:clrMap")):
            return True
        p = p.getparent()
    return False


def collect_style(prs, theme_colors, fonts):
    fill = collections.Counter()
    text = collections.Counter()
    ea, latin, heading = collections.Counter(), collections.Counter(), collections.Counter()
    geoms = collections.Counter()
    parts = [s for s in prs.slides] + list(prs.slide_layouts) + [prs.slide_master]
    for k, part in enumerate(parts):
        el = part._element.find(qn("p:cSld"))
        w = 3 if k < len(prs.slides) else 1  # 页面上的用色权重更高
        for clr in el.iter(qn("a:srgbClr"), qn("a:schemeClr")):
            if _skip(clr):
                continue
            v = clr.get("val")
            if clr.tag == qn("a:schemeClr"):
                v = theme_colors.get(SCHEME_ALIAS.get(v, v))
                if not v:
                    continue
            parent = clr.getparent()
            in_text = parent is not None and parent.getparent() is not None and \
                parent.getparent().tag in (qn("a:rPr"), qn("a:defRPr"))
            if in_text:
                if k < len(prs.slides):
                    text[norm_hex(v)] += 1
                continue
            # 填充色按形状面积加权：大面积的品牌色比小装饰更能代表模板
            sp = clr.getparent()
            while sp is not None and sp.tag not in (qn("p:sp"), qn("p:cxnSp")):
                sp = sp.getparent()
            area = 0.0
            if sp is not None:
                ext = sp.find(".//" + qn("a:ext"))
                if ext is not None and ext.get("cx"):
                    area = int(ext.get("cx")) * int(ext.get("cy")) / (prs.slide_width * prs.slide_height)
            # 渐变遮罩、半透明叠色只是氛围，不代表品牌色
            soft = clr.find(qn("a:alpha")) is not None or any(a.tag == qn("a:gradFill") for a in clr.iterancestors())
            fill[norm_hex(v)] += w * (1 + 20 * min(area, 0.6)) * (0.15 if soft else 1)
        for r in el.iter(qn("a:rPr"), qn("a:defRPr")):
            if _skip(r):
                continue
            for tag, ctr in ((qn("a:ea"), ea), (qn("a:latin"), latin)):
                f = r.find(tag)
                if f is not None and f.get("typeface"):
                    ctr[resolve_font(f.get("typeface"), fonts)] += 1
            f = r.find(qn("a:ea")) if r.find(qn("a:ea")) is not None else r.find(qn("a:latin"))
            if f is not None and r.get("sz") and int(r.get("sz")) >= 2800:
                heading[resolve_font(f.get("typeface"), fonts)] += 1
        for g in el.iter(qn("a:prstGeom")):
            geoms[g.get("prst")] += 1
    return fill, text, ea, latin, heading, geoms


def brand_dark(c):
    """海军蓝、墨绿这类深色品牌色：亮度低但有明显色相"""
    h, l, s_ = hls(c)
    return 0.07 < l <= 0.18 and s_ > 0.45


def pick_tokens(theme_colors, fonts, style):
    fill, text, ea, latin, heading, geoms = style
    both = fill + text
    sat = [c for c, _ in both.most_common() if is_saturated(c) or brand_dark(c)]
    for k in ("accent1", "accent2"):
        if not sat and theme_colors.get(k):
            sat.append(theme_colors[k])
    primary = sat[0] if sat else "2F5BEA"
    secondary = next((c for c in sat[1:] if hue_dist(c, primary) > 18), None) or mix(primary, "FFFFFF", 0.35)
    from common import contrast
    if contrast(primary, "FFFFFF") < 3 and contrast(secondary, "FFFFFF") >= 4.5 and \
            (is_saturated(secondary) or brand_dark(secondary)):
        primary, secondary = secondary, primary  # 金色这类浅强调色不适合当主色画卡片/图表
    darks = [c for c, _ in text.most_common() if luminance(c) < 0.06]
    dark = darks[0] if darks else (theme_colors.get("dk2") if luminance(theme_colors.get("dk2", "FFFFFF")) < 0.06
                                   else "1F2329")
    font_body = (ea.most_common(1)[0][0] if ea else None) or fonts["minorFont"]["ea"] or \
        (latin.most_common(1)[0][0] if latin else fonts["minorFont"]["latin"])
    font_head = heading.most_common(1)[0][0] if heading else (fonts["majorFont"]["ea"] or font_body)
    rounded = geoms.get("roundRect", 0) + geoms.get("ellipse", 0)
    square = geoms.get("rect", 0)
    return {
        "primary": primary, "secondary": secondary, "dark": dark,
        "body": mix(dark, "FFFFFF", 0.18), "muted": mix(dark, "FFFFFF", 0.45),
        "font": font_body, "font_heading": font_head,
        "font_latin": latin.most_common(1)[0][0] if latin else fonts["minorFont"]["latin"],
        "radius": 0.12 if rounded > square * 0.5 else 0.0,
        "palette": [c for c, _ in both.most_common(12)],
    }


# ---------- 每页分析 ----------
def ph_size(shape, slide):
    """占位符没有显式字号时，去版式和母版里找"""
    s = first_size(shape, 0)
    if s:
        return s
    try:
        idx = shape.placeholder_format.idx
        for lp in slide.slide_layout.placeholders:
            if lp.placeholder_format.idx == idx:
                s = first_size(lp, 0)
                if s:
                    return s
        typ = shape.placeholder_format.type
        tx = slide.slide_layout.slide_master._element.find(qn("p:txStyles"))
        st = tx.find(qn("p:titleStyle") if typ in TITLE_TYPES else qn("p:bodyStyle"))
        d = st.find(".//" + qn("a:defRPr"))
        if d is not None and d.get("sz"):
            return int(d.get("sz")) / 100
    except Exception:
        pass
    return 18.0


def is_number(t):
    return bool(t) and len(t) <= 12 and bool(NUM_RE.fullmatch(t.strip()))


def info(shape, box, slide, nested):
    t = shape_text(shape)
    ph = shape.is_placeholder
    size = ph_size(shape, slide) if ph else (max_size(shape) if t else first_size(shape))
    algn = None
    p = shape._element.find(".//" + qn("a:pPr"))
    if p is not None:
        algn = p.get("algn")
    return {
        "key": f"#{shape.shape_id}", "name": shape.name, "text": t, "box": [round(v, 2) for v in box],
        "size": round(size, 1), "wrap": not wrap_none(shape), "align": algn or "l",
        "placeholder": str(shape.placeholder_format.type).split(".")[-1].split(" ")[0] if ph else None,
        "nested": nested, "shape": shape,
    }


def classify(i, first, last, texts, W, H):
    """first / last：第一张和最后一张没隐藏的页（隐藏的「使用说明」页不算）"""
    all_text = "\n".join(t["text"] for t in texts)
    nums = [t for t in texts if is_number(t["text"])]
    if TOC_RE.search(all_text) or (i - first <= 2 and len(nums) >= 3):
        return "toc"
    if i >= last - 1 and i > first and END_RE.search(all_text):
        return "ending"
    disclaimer = any(DISCLAIMER_RE.search(t["text"]) and t["size"] <= 12 for t in texts)  # 免责声明都是小字
    if i == last and last - first >= 3 and (sum(len(t["text"]) for t in texts) < 40 or disclaimer):
        return "ending"  # 只有 logo / 口号，或者是免责声明的封底
    if i == first:
        return "cover"
    titles = [t for t in texts if not is_number(t["text"])]
    big = max(titles, key=lambda t: t["size"], default=None)
    big_txt = max((t for t in titles if t["text"].strip()), key=lambda t: t["size"], default=None)
    if any(QUOTE_MARK.fullmatch(t["text"]) for t in texts if t["text"]) or \
            (big_txt and big_txt["size"] >= 28 and big_txt["text"].lstrip().startswith(("“", "「", "『", '"'))):
        return "quote"
    if any(t["placeholder"] in ("TITLE", "CENTER_TITLE") for t in texts) and \
            any(t["placeholder"] == "TITLE" and t["box"][1] < H * 0.3 for t in texts):
        return "content"
    has_num = any(is_number(t["text"]) for t in texts)
    chars = sum(len(t["text"]) for t in texts)
    if big and big["box"][1] < H * 0.22 and big["size"] < 34:
        return "content"
    if (has_num or (big and big["size"] >= 28)) and chars < 80:
        return "section"
    return "content"


def reading_order(items):
    return sorted(items, key=lambda t: (round((t["box"][1] + t["box"][3] / 2) / 0.45), t["box"][0]))


def assign_slots(role, texts, W, H):
    slots = {}
    nums = [t for t in texts if is_number(t["text"])]
    words = [t for t in texts if not is_number(t["text"])]
    by_size = sorted(words, key=lambda t: (-t["size"], -t["box"][2] * t["box"][3]))
    title_ph = [t for t in words if t["placeholder"] in ("TITLE", "CENTER_TITLE")]
    if role == "toc":
        heads = [t for t in words if TOC_RE.search(t["text"])]
        rest = [t for t in words if t not in heads]
        # 目录条目字号相同：取出现最多的字号为条目，其余（如「目录」旁的英文副标）归入标题
        common = collections.Counter(t["size"] for t in rest).most_common(1)
        items = reading_order([t for t in rest if common and t["size"] == common[0][0]])
        others = [t for t in rest if t not in items]
        subs = []
        for it in items:  # 条目正下方、字号更小的文字 = 条目副标题
            b = it["box"]
            cand = [t for t in others if t["size"] < it["size"] and -0.15 <= t["box"][1] - (b[1] + b[3]) < 0.6
                    and abs(t["box"][0] - b[0]) < 0.3]
            subs.append(cand[0] if cand else None)
        if any(subs):
            slots["item_subs"] = [t["key"] if t else None for t in subs]
            others = [t for t in others if t not in subs]
        if len(items) <= 1:  # 一个文本框里写了好几段议题
            multi = next((t for t in rest if nav.para_items(t)), None)
            if multi:
                items, others = [multi], [t for t in rest if t is not multi]
                slots.pop("item_subs", None)
        heads += others
        slots["heading"] = [t["key"] for t in heads]
        slots["items"] = [t["key"] for t in items]
        slots["numbers"] = [t["key"] for t in reading_order(nums)]
    elif role == "quote":
        marks = [t for t in words if QUOTE_MARK.fullmatch(t["text"])]
        body = [t for t in by_size if t not in marks]
        if body:
            slots["title"] = body[0]["key"]
        slots["subtitle"] = [t["key"] for t in body[1:] if ATTRIB.match(t["text"])]
    elif role in ("cover", "ending", "section"):
        main = title_ph[0] if title_ph else (by_size[0] if by_size else None)
        if main:
            slots["title"] = main["key"]
        rest = [t for t in by_size if t is not main]
        from build import SAMPLE_RE
        deco = [t for t in rest if CAPS.fullmatch(t["text"].strip()) and main and t["size"] < main["size"]
                and not SAMPLE_RE.search(t["text"]) and not DATE_RE.search(t["text"])
                and not re.search(r"\d{4}[.\-/]\d{1,2}|(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC)\w*\s+\d",
                                  t["text"], re.I)]
        rest = [t for t in rest if t not in deco]
        dates = [t for t in rest if DATE_RE.search(t["text"])]
        slots["subtitle"] = [t["key"] for t in rest if t not in dates]
        if dates:
            slots["meta"] = [t["key"] for t in dates]
        if deco:
            slots["kicker"] = [t["key"] for t in deco]
        if nums:
            slots["number"] = nums[0]["key"]
    else:
        top = [t for t in words if t["box"][1] < H * 0.3]
        main = (title_ph[0] if title_ph else None) or (max(top, key=lambda t: (t["size"], t["box"][2]))
                                                         if top else None)
        if main:
            slots["title"] = main["key"]
        mb = main["box"][1] + main["box"][3] if main else H * 0.3
        hn = [t for t in nums if t["box"][1] + t["box"][3] <= mb + 0.1]  # 序号要在标题这一带，不是正文里的「1」
        if hn:
            slots["number"] = hn[0]["key"]
        if main:  # 标题旁边的小字（英文眉题、单位行）；正文占位符、大文本框不算
            slots["kicker"] = [t["key"] for t in top if t is not main and t["box"][1] < mb + 0.3
                               and t["box"][3] < H * 0.15 and t["placeholder"] not in ("BODY", "OBJECT", "PICTURE")
                               and abs(t["box"][0] - main["box"][0]) < 1.0]  # 和标题左对齐，不是图里的坐标轴标签
    return slots


def header_footer(slide_shapes, slots, W, H, layout_boxes):
    title = next((s for s in slide_shapes if s["key"] == slots.get("title")), None)
    head_keys = {slots.get("title"), slots.get("number")} | set(slots.get("kicker", []))
    if title:
        tb = title["box"]
        cut = tb[1] + tb[3]
        # 与标题同一水平带内的装饰（色块、线条、编号）都算页眉
        for s in slide_shapes:
            b = s["box"]
            if s["key"] in head_keys or (b[1] < cut + 0.15 and b[3] < H * 0.3 and b[1] + b[3] < H * 0.32):
                cut = max(cut, b[1] + b[3])
    else:
        cut = H * 0.18
    foot = H - 0.4
    for b, lim in [(s["box"], 0.88) for s in slide_shapes] + [(b, 0.84) for b in layout_boxes]:
        if b[1] > H * lim and b[3] < H * 0.16 and b[2] < W * 0.98:  # 页面上的形状要更贴底边才算页脚
            foot = min(foot, b[1] - 0.12)
    return cut + 0.25, foot


def layout_boxes(slide, W, H):
    """版式和母版上的非占位符形状（装饰、logo、页脚），整页大的背景图除外"""
    boxes = []
    layout = slide.slide_layout
    srcs = [layout.shapes]
    if layout._element.get("showMasterSp") != "0":
        srcs.append(layout.slide_master.shapes)
    for shapes in srcs:
        for s, b, _ in iter_shapes(shapes):
            if s.is_placeholder or b[2] * b[3] > W * H * 0.6 or s.shape_type == 6:
                continue
            boxes.append(b)
    return boxes


def free_rect(occupied, region, step=0.1):
    """region 内不与 occupied 相交的最大矩形（直方图法），返回 (x,y,w,h)"""
    x0, y0, x1, y1 = region
    cols, rows = max(1, int((x1 - x0) / step)), max(1, int((y1 - y0) / step))
    grid = [[True] * cols for _ in range(rows)]
    for (bx, by, bw, bh) in occupied:
        c0, c1 = int((bx - x0) / step), int((bx + bw - x0) / step + 0.999)
        r0, r1 = int((by - y0) / step), int((by + bh - y0) / step + 0.999)
        for r in range(max(0, r0), min(rows, r1)):
            for c in range(max(0, c0), min(cols, c1)):
                grid[r][c] = False
    return _max_rect(grid, x0, y0, step)


def _max_rect(grid, x0, y0, step):
    rows, cols = len(grid), len(grid[0])
    hist = [0] * cols
    best = (0, 0, 0, 0, 0)
    for r in range(rows):
        for c in range(cols):
            hist[c] = hist[c] + 1 if grid[r][c] else 0
        stack = []
        for c in range(cols + 1):
            h = hist[c] if c < cols else 0
            start = c
            while stack and stack[-1][1] >= h:
                s, sh = stack.pop()
                area = sh * (c - s)
                if area > best[0]:
                    best = (area, s, r - sh + 1, c - s, sh)
                start = s
            stack.append((start, h))
    _, c, r, w, h = best
    return (round(x0 + c * step, 2), round(y0 + r * step, 2), round(w * step, 2), round(h * step, 2))


def busy_boxes(img_path, W, H, region, step=0.1):
    """从渲染图里找出背景上的「画」（吉祥物、插画、山水），作为占用区域。
    按 0.1 英寸格子统计边缘密度：细线、网格、纹理密度低不算占用，大块图形和文字密度高才算。"""
    im = Image.open(img_path).convert("L")
    edges = im.filter(ImageFilter.FIND_EDGES).point(lambda v: 255 if v > 24 else 0)
    cols, rows = int(W / step), int(H / step)
    blk = 3  # 边缘密度按 0.3 英寸的大格子算：细线、网格占比低，文字和插画占比高
    coarse = edges.resize((max(1, cols // blk), max(1, rows // blk)), Image.BOX)
    small = coarse.resize((cols, rows), Image.NEAREST)
    # 大块平涂图形（山体、色块插画）边缘少：和「同一行、同一列的中位色」都差得多才算（渐变背景不会误报）
    col = Image.open(img_path).convert("RGB").resize((cols, rows), Image.BOX)
    cp = col.load()

    def med(vals):
        return tuple(sorted(v[k] for v in vals)[len(vals) // 2] for k in range(3))

    rmed = [med([cp[c, r] for c in range(cols)]) for r in range(rows)]
    cmed = [med([cp[c, r] for r in range(rows)]) for c in range(cols)]

    def dist(p, q):
        return sum(abs(p[k] - q[k]) for k in range(3))

    px = small.load()
    fine = edges.resize((cols, rows), Image.BOX).load()
    busy = {(c, r) for r in range(rows) for c in range(cols) if px[c, r] > 64}
    diff = {(c, r) for r in range(rows) for c in range(cols)
            if (c, r) not in busy and min(dist(cp[c, r], rmed[r]), dist(cp[c, r], cmed[c])) > 30}
    busy |= solid_regions(diff, fine, cols, rows)
    grown = {(c + dc, r + dr) for c, r in busy for dc in (-1, 0, 1) for dr in (-1, 0, 1)}
    x0, y0, x1, y1 = region
    return [(c * step, r * step, step, step) for c, r in grown
            if 0 <= c < cols and 0 <= r < rows and x0 <= c * step < x1 and y0 <= r * step < y1]


def solid_regions(diff, fine, cols, rows):
    """和底色差得多的连片格子里，挑出「边界清楚」的：山体、色块插画的轮廓是一条清楚的线；
    背景上的柔光、光晕边界发虚，不算占用（否则深色科技风模板的内容区会被砍掉一大块）"""
    out, seen = set(), set()
    for start in diff:
        if start in seen:
            continue
        comp, stack = [], [start]
        seen.add(start)
        while stack:
            c, r = stack.pop()
            comp.append((c, r))
            for nb in ((c + 1, r), (c - 1, r), (c, r + 1), (c, r - 1)):
                if nb in diff and nb not in seen:
                    seen.add(nb)
                    stack.append(nb)
        # 区域的边界（页面边缘不算：渲染图的四边本身就有一圈假边缘）
        border = [(c, r) for c, r in comp if 0 < c < cols - 1 and 0 < r < rows - 1
                  and any(nb not in diff for nb in ((c + 1, r), (c - 1, r), (c, r + 1), (c, r - 1)))]
        sharp = sum(1 for c, r in border if fine[c, r] > 10)
        if len(comp) <= 4 or (border and sharp >= 0.25 * len(border)):
            out.update(comp)
    return out


def sample_bg(img_path, W, H, rect):
    im = Image.open(img_path).convert("RGB")
    sx, sy = im.width / W, im.height / H
    x, y, w, h = rect
    x0, y0 = int(x * sx), int(y * sy)
    crop = im.crop((x0, y0, max(x0 + 1, int((x + w) * sx)), max(y0 + 1, int((y + h) * sy)))).resize((40, 24))
    c = collections.Counter(crop.getdata()).most_common(1)[0][0]
    return to_hex(c)


def strip_samples(prs, plan):
    """删掉内容页的示例内容，供渲染「干净」的版面"""
    for i, slide in enumerate(prs.slides):
        for key in plan.get(i, []):
            for s in list(slide.shapes):
                if f"#{s.shape_id}" == key:
                    s._element.getparent().remove(s._element)


def is_footer(s):
    """页脚、页码、日期占位符，以及含页码域的文本框：模板自己维护，不当槽位"""
    if s.is_placeholder:
        try:
            if s.placeholder_format.type in FOOTER_TYPES:
                return True
        except ValueError:
            pass
    return s._element.find(".//" + qn("a:fld")) is not None


def toc_mode(entry, texts, shapes):
    """目录是「一个文本框多段」还是「每条一个形状」，以及有没有高亮当前项（议程式章节页）"""
    items = entry["slots"].get("items", [])
    by_key = {t["key"]: t for t in texts}
    pi = nav.para_items(by_key[items[0]]) if len(items) == 1 and items[0] in by_key else None
    if pi:
        entry["toc_mode"] = "paragraphs"
        entry["toc_current"] = pi["current"]
        entry["toc_mark"] = nav.para_mark(by_key[items[0]], shapes, pi["current"])
    elif len(items) >= 3:
        entry["toc_current"] = nav.highlight([nav.rpr_sig(nav.first_rpr(by_key[k]["shape"]._element))
                                              for k in items if k in by_key])


def refine_roles(entries, W, H):
    """看全局再改角色：
    - 有带序号的章节页时，没序号、标题在左半边的「章节页」其实是「左标题、右内容」的内容页
    - 有两张以上目录页时，高亮了某一项的那几张是议程式章节页"""
    secs = [e for e in entries if e["role"] == "section"]
    if any(e["slots"].get("number") for e in secs):
        for e in secs:
            title = next((t for t in e["texts"] if t["key"] == e["slots"].get("title")), None)
            if not e["slots"].get("number") and title and title["box"][0] + title["box"][2] < W * 0.55 \
                    and title["box"][1] > H * 0.25:
                e["role"], e["side_title"] = "content", True
    tocs = [e for e in entries if e["role"] == "toc"]
    lit = [e for e in tocs if e.get("toc_current") is not None]
    # 有一张没高亮的总目录，或者几张目录高亮的位置不同，才说明高亮是在标「当前章」
    if len(tocs) >= 2 and lit and (len(lit) < len(tocs) or len({e["toc_current"] for e in lit}) > 1):
        for e in lit:
            e["role"], e["agenda"] = "section", True


def is_backdrop(d, W, H):
    """整页的背景图、底色块（不是图表和表格：示例图表常常也占大半页）"""
    a = d["box"][2] * d["box"][3]
    return a > W * H * 0.8 or (a > W * H * 0.6 and d["shape"].shape_type in (1, 13))


def content_region(slide, shapes, entry, W, H):
    """内容页：版面区域、要删的示例内容、被占用的地方"""
    slots = entry["slots"]
    lb = layout_boxes(slide, W, H)
    trk = entry.get("tracker") or {}
    keep = set(trk.get("items", [])) | {trk.get("mark")}
    samples = []
    if entry.get("side_title"):  # 标题在左侧面板里：整页去掉占用的地方，剩下的就是内容区
        foot = header_footer(shapes, {}, W, H, lb)[1]
        region = (0.45, 0.45, W - 0.45, foot)
    else:
        top, foot = header_footer(shapes, slots, W, H, lb)
        left = min([shapes_by(shapes, slots.get("title"))["box"][0]] if slots.get("title") else [0.6])
        left = max(0.45, min(left, 1.2))
        region = (left, top, W - left, foot)
        for d in shapes:
            if d["nested"] or d["key"] in keep:
                continue
            b = d["box"]
            if is_backdrop(d, W, H):
                continue
            empty_ph = d["shape"].is_placeholder and not d["text"] and d["key"] != slots.get("title") \
                and not is_footer(d["shape"])
            inside = empty_ph or b[1] >= top - 0.3 and b[1] + b[3] <= foot + 0.3 and (b[0] >= left - 0.8 or b[0] + b[2] / 2 >= left)
            edge_pic = d["shape"].shape_type == 13 and (b[0] < 0.1 or b[0] + b[2] > W - 0.1 or b[1] + b[3] > H - 0.1)
            if inside and not edge_pic and not is_footer(d["shape"]) and \
                    d["key"] not in (slots.get("title"), slots.get("number")):
                samples.append(d["key"])
    occ = lb + [d["box"] for d in shapes if not d["nested"] and d["key"] not in samples and not is_backdrop(d, W, H)]
    return {"region": [round(v, 2) for v in region], "samples": samples, "_occ": occ}


def analyze(path, out, do_render=True):
    prs = Presentation(path)
    W, H = inch(prs.slide_width), inch(prs.slide_height)
    theme_colors, fonts = theme_info(prs)
    tokens = pick_tokens(theme_colors, fonts, collect_style(prs, theme_colors, fonts))
    slides = []
    n = len(prs.slides)
    # 在 ≥3 页上重复出现的文字是品牌名/页脚（如「睿信资本」），保留原样，不当作可填槽位
    seen = collections.Counter()
    for slide in prs.slides:
        seen.update({shape_text(s) for s, _, _ in iter_shapes(slide.shapes) if s.has_text_frame and shape_text(s)})
    from build import SAMPLE_RE
    brand = {t for t, c in seen.items() if c >= 3 and len(t) <= 30 and not NUM_RE.fullmatch(t.strip())
             and not SAMPLE_RE.search(t)}
    hidden = [sl._element.get("show") == "0" for sl in prs.slides]
    visible = [i for i in range(n) if not hidden[i]] or [0]
    first, last = visible[0], visible[-1]
    pages = []
    for i, slide in enumerate(prs.slides):  # 第一遍：每页的角色和槽位
        shapes, texts = [], []
        for s, b, nested in iter_shapes(slide.shapes):
            d = info(s, b, slide, nested)
            shapes.append(d)
            if s.has_text_frame and (d["text"] or s.is_placeholder) and not is_footer(s):
                texts.append(d)
        role = "hidden" if hidden[i] else classify(i, first, last, texts, W, H)
        ids = collections.Counter(d["key"] for d in shapes)
        entry = {"index": i, "role": role, "layout": slide.slide_layout.name,
                 "texts": [{k: v for k, v in t.items() if k != "shape"} for t in texts]}
        if any(c > 1 for c in ids.values()):
            entry["dup_ids"] = sorted(k for k, c in ids.items() if c > 1)
        trk = nav.detect_tracker(texts, shapes, W, H) if role == "content" else None
        if trk:
            entry["tracker"] = trk
            texts = [t for t in texts if t["key"] not in trk["items"]]
        entry["slots"] = assign_slots(role, [t for t in texts if t["text"] not in brand], W, H)
        if role == "toc":
            toc_mode(entry, texts, shapes)
        pages.append((slide, shapes, texts, entry))
    refine_roles([e for *_, e in pages], W, H)
    for slide, shapes, texts, entry in pages:  # 第二遍：内容页的安全区和示例内容
        if entry["role"] == "content":
            entry.update(content_region(slide, shapes, entry, W, H))
        slides.append(entry)

    files = []
    if do_render:
        from render import render, sheets
        stripped = os.path.join(out, "_stripped.pptx")
        p2 = Presentation(path)
        strip_samples(p2, {s["index"]: s.get("samples", []) for s in slides})
        for sl in p2.slides:  # 隐藏页 LibreOffice 不导出，会让图和页错位
            sl._element.attrib.pop("show", None)
        p2.save(stripped)
        files = render(stripped, os.path.join(out, "template-render"), dpi=60)
        os.remove(stripped)
        if len(files) != n:
            print(f"⚠️ 渲染出 {len(files)} 页、模板有 {n} 页，跳过像素分析", file=sys.stderr)
            files = []
    for s in slides:
        img = files[s["index"]] if files else None
        if s["role"] == "content":
            reg = s["region"]
            occ = s.pop("_occ")
            if img:
                occ = occ + busy_boxes(img, W, H, (reg[0], reg[1], reg[2], reg[3]))
            a = free_rect(occ, reg)
            if a[2] < 0.3 or a[3] < 0.3:  # 整页都是图：退回到几何区域，交给模型判断
                a = (reg[0], reg[1], reg[2] - reg[0], reg[3] - reg[1])
                s["warning"] = "背景图占满了内容区，安全区按几何位置给出，请看渲染图确认"
            pad = 0.05
            s["area"] = [round(a[0] + pad, 2), round(a[1], 2), round(a[2] - 2 * pad, 2), round(a[3], 2)]
            if (a[2] < W * 0.45 or a[3] < H * 0.35) and "warning" not in s:
                s["warning"] = "内容安全区偏小，这一页不适合放大段内容"
        rect = s.get("area") or [W * 0.3, H * 0.3, W * 0.4, H * 0.4]
        if img:
            s["bg"] = sample_bg(img, W, H, rect)
        else:
            s["bg"] = "FFFFFF"
        s["dark"] = luminance(s["bg"]) < 0.22
    prof = {"template": os.path.relpath(os.path.abspath(path), os.path.abspath(out)), "size": [round(W, 3), round(H, 3)], "tokens": tokens,
            "theme_colors": theme_colors, "theme_fonts": fonts, "slides": slides}
    content_bgs = [s["bg"] for s in slides if s["role"] == "content" and not s["dark"]]
    prof["tokens"]["surface"] = "FFFFFF" if not content_bgs or luminance(content_bgs[0]) < 0.9 else \
        mix(content_bgs[0], tokens["primary"], 0.04)
    if files:
        labels = [f'{s["role"]}' for s in slides]
        prof["sheets"] = [os.path.relpath(p, out) for p in
                          sheets(files, os.path.join(out, "template-render"), labels=labels)]
    return prof


def shapes_by(shapes, key):
    return next(s for s in shapes if s["key"] == key)


def capacity(t, W):
    """估算一个槽位一行能放多少个汉字"""
    w = t["box"][2] if t["wrap"] else W * 0.8
    return max(1, int(w * 72 / max(6, t["size"])))


def write_md(prof, path):
    W, H = prof["size"]
    tk = prof["tokens"]
    L = [f"# 模板画像：{os.path.basename(prof['template'])}", "",
         f"- 页面：{W:.2f} × {H:.2f} 英寸，共 {len(prof['slides'])} 页",
         f"- 主色 primary `#{tk['primary']}`，辅色 secondary `#{tk['secondary']}`，标题文字 dark `#{tk['dark']}`，"
         f"正文 body `#{tk['body']}`，弱化 muted `#{tk['muted']}`，卡片底 surface `#{tk['surface']}`",
         f"- 字体：正文 `{tk['font']}`，标题 `{tk['font_heading']}`，西文 `{tk['font_latin']}`；"
         f"圆角 radius={tk['radius']}（0 = 直角风格）",
         f"- 模板高频色：{' '.join('#' + c for c in tk['palette'])}", ""]
    if prof.get("sheets"):
        L += ["模板缩略图（每格右下角是页码 · 角色，路径相对本文件）：" + "，".join(prof["sheets"]), ""]
    by_role = collections.defaultdict(list)
    for s in prof["slides"]:
        by_role[s["role"]].append(s["index"])
    L += ["## 角色总览（deck.json 的 from 填这里的页码，从 0 开始）", ""]
    for r in ("cover", "toc", "section", "content", "quote", "ending", "hidden", "other"):
        if by_role.get(r):
            L.append(f"- **{r}**：{by_role[r]}")
    L.append("")
    for s in prof["slides"]:
        L.append(f"## 第 {s['index']} 页 · {s['role']}（版式：{s['layout']}，底色 #{s['bg']}"
                 f"{'，深色底' if s['dark'] else ''}）")
        for k, v in s["slots"].items():
            L.append(f"- 槽位 `{k}` → {v}")
        if s.get("area"):
            a = s["area"]
            L.append(f"- 内容安全区 area = [{a[0]}, {a[1]}, {a[2]}, {a[3]}]（x, y, 宽, 高，英寸）")
        if s["role"] == "hidden":
            L.append("- 隐藏页（模板使用说明之类），不用来做内容")
        if s["role"] == "quote":
            L.append("- 模板自带的金句页：`title` 写金句，`subtitle` 写「—— 出处」")
        if s.get("agenda"):
            L.append(f"- 议程式章节页：写 `items`（全部章节名）+ `current`（当前章，从 0 开始），"
                     f"模板高亮第 {s['toc_current']} 项的样子会移到 current")
        elif s.get("toc_mode") == "paragraphs":
            L.append("- 目录是一个文本框里一段一个议题：`items` 写章节名，编号格式沿用模板")
        if s.get("dup_ids"):
            L.append(f"- ⚠️ 这一页有重复的形状 id {s['dup_ids']}：按 `#id` 引用时可能找错形状，最好在 PowerPoint 里另存一次让它重新编号")
        if s.get("tracker"):
            t = s["tracker"]
            L.append(f"- 进度标签（{'一排' if t['axis'] == 'row' else '一列'}，{len(t['items'])} 格 {t['items']}，模板高亮第 {t['current']} 格"
                     f"{'，标记 ' + t['mark'] if t.get('mark') else ''}）：写 `tracker: {{items: [短章节名], current: 当前章}}`；"
                     "没有章节就写 `tracker: null` 删掉")
        if s.get("side_title"):
            L.append("- 标题在左侧面板里，内容放右边（area 已算好）")
        if s.get("samples"):
            L.append(f"- 构建时自动删除的示例内容：{', '.join(s['samples'])}")
        if s.get("warning"):
            L.append(f"- ⚠️ {s['warning']}")
        L.append("")
        L.append("| 形状 | 文字 | 字号 | 换行 | 一行约容纳 |")
        L.append("|---|---|---|---|---|")
        for t in s["texts"]:
            txt = (t["text"] or "（空占位符）").replace("\n", " / ").replace("|", "｜")
            txt = txt[:40] + ("…" if len(txt) > 40 else "")
            L.append(f"| `{t['key']}` {t['name']} | {txt} | {t['size']} | {'是' if t['wrap'] else '否（构建时自动加宽）'} "
                     f"| {capacity(t, W)} 字 |")
        L.append("")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(L))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("template")
    ap.add_argument("--out", default="work")
    ap.add_argument("--no-render", action="store_true")
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    prof = analyze(a.template, a.out, not a.no_render)
    save_json(os.path.join(a.out, "profile.json"), prof)
    write_md(prof, os.path.join(a.out, "profile.md"))
    roles = collections.Counter(s["role"] for s in prof["slides"])
    print(f"模板 {len(prof['slides'])} 页：" + "，".join(f"{k}×{v}" for k, v in roles.items()))
    print(f"主色 #{prof['tokens']['primary']}  字体 {prof['tokens']['font']}")
    print("→", os.path.join(a.out, "profile.md"))
    if not any(s["role"] == "content" for s in prof["slides"]):
        print("⚠️ 没识别出内容页：请在 deck.json 里用 area 手动指定内容区域", file=sys.stderr)


if __name__ == "__main__":
    main()
