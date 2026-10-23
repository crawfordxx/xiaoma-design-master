"""按 deck.json 用模板生成 pptx。用法：
    python3 build.py deck.json [--out out/deck.pptx]

每一页都从模板的某一页复制出来（背景、母版、logo、装饰原样保留），再：
1. 往模板槽位里填字：保留原字体字号颜色；不换行的框自动加宽，放不下才缩字；
2. 删掉模板示例内容（profile 里的 samples），以及没填的「标题/XXXX/请输入」这类示例字；
3. 在内容安全区里按 body 画版式（recipes.py）；
4. 写演讲者备注。
"""
import argparse
import copy
import os
import re
import sys

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

import draw
import nav
import recipes
from common import emu, inch, iter_shapes, load_json, overlap, save_json, text_em, text_insets, wrap_none, count_lines
from draw import Theme

SAMPLE_RE = re.compile(r"X{3,}|x{4,}|^\s*(大?标题|副标题|小标题|正文|内容|content|title|subtitle|text)\s*$|"
                       r"^.{0,6}(标题|副标题|正文|内容|说明|文字|描述|简介)(内容|文字)?\s*$|\bSUB-?TITLE\b|"
                       r"(请|单击|点击|在此|此处)(输入|添加|编辑|替换)|click to|lorem|your (title|text)|"
                       r"标题文字|添加标题|输入标题|关键词|20XX|\bXX\b|^\s*0?\d\s*$", re.I)
REL_ATTRS = [qn("r:embed"), qn("r:id"), qn("r:link"), qn("r:pict"), qn("r:dm"), qn("r:lo"), qn("r:qs"), qn("r:cs")]
CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")  # 从 Word 粘贴来的控制字符，写进 XML 会报错


# ---------- 复制页面 ----------
def clone(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid = {}
    for rel in src.part.rels.values():
        if rel.reltype in (RT.NOTES_SLIDE, RT.SLIDE_LAYOUT):
            continue
        if rel.is_external:
            rid[rel.rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid[rel.rId] = new.part.relate_to(rel.target_part, rel.reltype)
    s_cSld, n_cSld = src._element.cSld, new._element.cSld
    bg = s_cSld.find(qn("p:bg"))
    if bg is not None:
        n_cSld.insert(0, copy.deepcopy(bg))
    tree = n_cSld.spTree  # slide.shapes 缓存了这个节点，只能换子节点不能换节点本身
    for ch in list(tree):
        tree.remove(ch)
    for ch in s_cSld.spTree:
        tree.append(copy.deepcopy(ch))
    for el in n_cSld.iter():
        for a in REL_ATTRS:
            v = el.get(a)
            if v and v in rid:
                el.set(a, rid[v])
    if src._element.get("showMasterSp") is not None:  # 「隐藏母版装饰」这类页面级属性
        new._element.set("showMasterSp", src._element.get("showMasterSp"))
    ovr = src._element.find(qn("p:clrMapOvr"))
    if ovr is not None:
        old = new._element.find(qn("p:clrMapOvr"))
        if old is not None:
            old.getparent().replace(old, copy.deepcopy(ovr))
        else:
            new._element.cSld.addnext(copy.deepcopy(ovr))
    for tag in ("p:transition", "p:timing"):
        t = src._element.find(qn(tag))
        if t is not None:
            new._element.append(copy.deepcopy(t))
    return new


def drop_originals(prs, n):
    lst = prs.slides._sldIdLst
    for sid in list(lst)[:n]:
        prs.part.drop_rel(sid.get(qn("r:id")))
        lst.remove(sid)


# ---------- 槽位 ----------
def find_shape(slide, key):
    for s, b, nested in iter_shapes(slide.shapes):
        if (key.startswith("#") and f"#{s.shape_id}" == key) or s.name == key:
            return s, b, nested
    return None, None, None


def set_lines(shape, lines):
    """逐段替换文字，每段保留模板第一段 run 的格式"""
    txBody = shape.text_frame._txBody
    ps = txBody.findall(qn("a:p"))
    proto = next((r for p in ps for r in p.findall(qn("a:r"))), None)
    for i, line in enumerate(lines):
        if i < len(ps):
            p = ps[i]
        else:
            p = copy.deepcopy(ps[-1])
            txBody.append(p)
            ps.append(p)
        runs = p.findall(qn("a:r"))
        if not runs:
            r = copy.deepcopy(proto) if proto is not None else None
            if r is None:
                shape.text_frame.paragraphs[i].text = line
                continue
            end = p.find(qn("a:endParaRPr"))
            if end is not None:
                end.addprevious(r)
            else:
                p.append(r)
            runs = [r]
        runs[0].find(qn("a:t")).text = CTRL.sub("", line)
        for r in runs[1:]:
            p.remove(r)
        for br in p.findall(qn("a:br")):
            p.remove(br)
    for p in ps[len(lines):]:
        txBody.remove(p)


def run_sizes(shape):
    out = []
    for el in shape._element.iter(qn("a:rPr"), qn("a:endParaRPr")):
        out.append(int(el.get("sz")) / 100 if el.get("sz") else None)
    return out


def scale_font(shape, k, base):
    for el in shape._element.iter(qn("a:rPr"), qn("a:endParaRPr")):
        sz = int(el.get("sz")) / 100 if el.get("sz") else base
        el.set("sz", str(int(max(8, sz * k) * 100)))


SAFETY = 1.1  # 渲染字体比估算宽一点，留余量


def resize(shape, box, nx, ny, nw, nh):
    """按绝对坐标改位置尺寸；组合内的形状要换算回组合的子坐标"""
    # 先把四个值都读出来再写：位置继承自版式的占位符，写了 left 以后再读 top 会读到 0
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    kx = box[2] / max(1e-6, inch(width)) if width else 1
    ky = box[3] / max(1e-6, inch(height)) if height else 1
    shape.left = left + emu((nx - box[0]) / kx)
    shape.top = top + emu((ny - box[1]) / ky)
    shape.width = emu(nw / kx)
    shape.height = emu(nh / ky)


def right_limit(slide, box, W):
    """同一水平带里，右边第一个挡路的形状（logo 等）的左边缘"""
    lim = W - 0.3
    for s, b, nested in iter_shapes(slide.shapes):
        if s.shape_type == 6 or b[0] <= box[0] + box[2] * 0.5 or b[2] * b[3] > 0.3 * W * W:
            continue
        if b[0] <= box[0] + 0.05 or overlap(b, box) >= 0.9 * b[2] * b[3]:
            continue
        if b[1] < box[1] + box[3] and b[1] + b[3] > box[1]:
            lim = min(lim, b[0] - 0.15)
    return lim


def push_above(slide, shape, box, dy):
    for s, b, nested in list(iter_shapes(slide.shapes)):
        if s is shape or nested or s.shape_type == 6 or not (s.has_text_frame and s.text_frame.text.strip()):
            continue
        if box[1] - 0.6 <= b[1] + b[3] <= box[1] + 0.05 and b[1] - dy > 0.15 and overlap(
                (b[0], 0, b[2], 1), (box[0], 0, box[2], 1)) > 0:
            resize(s, b, b[0], b[1] - dy, b[2], b[3])  # 四个值一起写：位置继承版式的占位符只改 top 会被挪到最左边


def below_neighbor(slide, shape, box):
    for s, b, nested in iter_shapes(slide.shapes):
        if s is shape or s.shape_type == 6 or not (s.has_text_frame and s.text_frame.text.strip()):
            continue
        if b[1] >= box[1] + box[3] * 0.6 and b[1] < box[1] + box[3] + 0.6 and overlap(
                (b[0], 0, b[2], 1), (box[0], 0, box[2], 1)) > 0:
            return True
    return False


def fit_slot(slide, shape, box, base_size, W):
    """填完字后让它放得下：不换行的框先加宽（保持对齐方式、不压到右边的 logo），多行时往上长，实在放不下才缩字"""
    lines = [p.text for p in shape.text_frame.paragraphs]
    size = max([s for s in run_sizes(shape) if s] or [base_size])
    l_in, r_in, t_in, b_in = text_insets(shape)
    x, y, w, h = box
    msg = None
    nowrap = wrap_none(shape)
    wrap2 = False
    need = max(text_em(t) for t in lines) * size / 72 * SAFETY + l_in + r_in
    if nowrap or (len(lines) == 1 and need > w):
        if need > w:
            p = shape._element.find(".//" + qn("a:pPr"))
            algn = (p.get("algn") if p is not None else None) or "l"
            lim_r = right_limit(slide, box, W)
            room = (lim_r - x) if algn == "l" else (W - 0.6)
            nw = min(need, room)
            nx = x - (nw - w) / 2 if algn == "ctr" else (x - (nw - w) if algn == "r" else x)
            nx = max(min(0.3, x), min(nx, W - min(0.3, W - x - w) - nw))
            resize(shape, box, nx, y, nw, h)
            box = (nx, y, nw, h)
            # 会换行的框：折成两行比缩成小字好（标题旁边常有产品图、logo 挡着加宽）
            wrap2 = not nowrap and len(lines) == 1 and need <= 2 * (nw - l_in - r_in) / SAFETY + l_in + r_in
            if wrap2 and need > nw + 0.01:  # 在空格 / 标点处断开，别把词拆散；没有好断点就缩字
                two = split_balanced(lines[0])
                if len(two) == 2:
                    lines = two
                    set_lines(shape, lines)
                elif (nw - l_in - r_in) / (need - l_in - r_in) >= 0.75:
                    wrap2 = False  # 缩到 75% 以上还能放下：缩字比随便断行好看
            if need > nw + 0.01 and not wrap2:
                k = max(0.5, (nw - l_in - r_in) / (need - l_in - r_in))
                scale_font(shape, k, base_size)
                size *= k
                msg = f"缩字 {k:.0%}"
    body_w = box[2] - l_in - r_in
    nl = len(lines) if nowrap else sum(count_lines(t, body_w, size) for t in lines)
    need_h = nl * size * 1.32 / 72 + t_in + b_in
    if need_h > box[3] + 0.05:
        if nl > 1 and len(lines) > 1 or nowrap or wrap2:
            # 多行标题：下面有字就往上长，否则往下长
            up = below_neighbor(slide, shape, box)
            ny = box[1] - (need_h - box[3]) if up else box[1]
            if up:  # 往上长时，紧贴在上面的眉题、小字一起往上推，别被压住
                push_above(slide, shape, box, box[1] - max(0.2, ny))
            resize(shape, box, box[0], max(0.2, ny), box[2], need_h)
        elif need_h > box[3] * 1.25 + 0.1:
            k = 1.0
            while k > 0.6 and nl * size * k * 1.2 / 72 + t_in + b_in > box[3] * 1.25 + 0.1:
                k -= 0.05
                nl = sum(count_lines(t, body_w, size * k) for t in lines)
            if k < 1:
                scale_font(shape, k, base_size)
                msg = f"缩字 {k:.0%}"
    return msg


def split_balanced(text):
    """一行标题折成两行：在最靠中间的空格、标点、破折号后断开；找不到合适的断点返回原文（交给缩字）"""
    cuts = [m.end() if m.group(0).strip() else m.start() for m in re.finditer(r"\s+|[，、：；,:;·]|——", text)]
    cuts = [c for c in cuts if 0.25 * len(text) < c < 0.75 * len(text)]
    if not cuts:
        return [text]
    c = min(cuts, key=lambda k: abs(text_em(text[:k]) - text_em(text[k:])))
    return [text[:c].strip(), text[c:].strip()]


def shape_in_group(shape):
    return shape._element.getparent().tag == qn("p:grpSp")


def remove(shape):
    el = shape._element
    el.getparent().remove(el)


def top_group(slide, shape):
    el = shape._element
    while el.getparent() is not None and el.getparent().tag == qn("p:grpSp"):
        el = el.getparent()
    return el


def fill_slot(slide, key, value, W, log):
    shape, box, _ = find_shape(slide, key)
    if shape is None:
        log.append(f"找不到槽位 {key}")
        return
    if value is None or value == "" or value == []:
        remove(shape)
        return
    lines = value if isinstance(value, list) else str(value).split("\n")
    base = next((s for s in run_sizes(shape) if s), 18)
    set_lines(shape, [str(v) for v in lines])
    r = fit_slot(slide, shape, box, base, W)
    if r:
        log.append(f"{key} {r}")


# ---------- 每页 ----------
ROLE_KEYS = ("title", "subtitle", "number", "meta", "kicker", "heading")


def apply_slots(slide, spec, prof_slide, W, log):
    slots = prof_slide["slots"]
    filled = set()
    for k in ROLE_KEYS:
        if k not in spec or k not in slots:
            if spec.get(k) is not None and k not in slots and k not in ("kicker", "heading"):
                log.append(f"模板第 {prof_slide['index']} 页没有 {k} 槽位，{k} 未使用")
            continue
        targets = slots[k] if isinstance(slots[k], list) else [slots[k]]
        if not targets and spec.get(k):
            log.append(f"模板第 {prof_slide['index']} 页的 {k} 槽位是空的，{k} 未使用")
        vals = spec[k] if isinstance(spec[k], list) and isinstance(slots[k], list) else [spec[k]]
        for i, key in enumerate(targets):
            if i < len(vals):
                if k == "title" and vals[i] is None:
                    drop_title_decor(slide, key)
                fill_slot(slide, key, vals[i], W, log)
                filled.add(key)
            # 没给值的槽位不在这里删：后面统一只删「示例字」，品牌名、口号这类模板原文保留
    if prof_slide["role"] == "section" and str(spec.get("number", "")).isdigit() and "kicker" not in spec:
        renumber_kickers(slide, slots.get("kicker", []), int(spec["number"]))
    if (prof_slide["role"] == "toc" or prof_slide.get("agenda")) and "items" in spec:
        fill_toc(slide, spec["items"], slots, W, log, prof_slide, spec.get("current"))
        filled |= set(slots.get("items", [])) | set(slots.get("numbers", [])) | set(slots.get("item_subs", []))
    trk = prof_slide.get("tracker")
    if trk:
        filled |= set(trk["items"])
        if "tracker" not in spec:
            log.append(f"模板这一页有进度标签 {trk['items']}，deck.json 没写 tracker：标签保留了模板原文")
        else:
            t = spec["tracker"] or {}
            nav.fill_tracker(slide, trk, [str(x) for x in t.get("items") or []], t.get("current"), find_shape, remove,
                             log)
    for k, v in (spec.get("slots") or {}).items():
        fill_slot(slide, k, v, W, log)
        filled.add(k)
    keep = set(spec.get("keep", []))
    for t in prof_slide["texts"]:
        if t["key"] in filled or t["key"] in keep:
            continue
        # 章节页、内容页上没填的副标题、日期位也是示例字（封面上常是装饰大字，不动）
        unfilled = prof_slide["role"] in ("section", "content", "quote") and any(
            t["key"] in (slots[k] if isinstance(slots[k], list) else [slots[k]]) for k in ("subtitle", "meta") if k in slots)
        if SAMPLE_RE.search(t["text"] or "") or not t["text"] or unfilled:
            shape, _, _ = find_shape(slide, t["key"])
            if shape is not None:
                remove(shape)


EN_NUM = ["ZERO", "ONE", "TWO", "THREE", "FOUR", "FIVE", "SIX", "SEVEN", "EIGHT", "NINE", "TEN", "ELEVEN", "TWELVE"]


def renumber_kickers(slide, keys, n):
    """章节页上整段是「CHAPTER ONE」「PART 01」这种章节序号的眉题，序号改成当前章（保留原大小写）"""
    words = "|".join(EN_NUM[1:])
    pat = re.compile(rf"^(\s*(?:PART|CHAPTER|SECTION|CH\.?|NO\.?)\s*)({words}|\d{{1,2}})(\s*)$", re.I)
    for key in keys:
        shape, _, _ = find_shape(slide, key)
        if shape is None:
            continue
        for t in shape._element.iter(qn("a:t")):
            m = pat.match(t.text or "")
            if not m:
                continue
            old = m.group(2)
            if old.isdigit():
                new = f"{n:0{len(old)}d}"
            else:
                new = EN_NUM[n] if n < len(EN_NUM) else str(n)
                new = new.lower() if old.islower() else (new.title() if old.istitle() else new)
            t.text = m.group(1) + new + m.group(3)


def drop_title_decor(slide, key):
    """删标题时，把紧贴标题的小装饰（竖条、短横线、序号底）一起删掉，免得留下孤零零的装饰"""
    shape, tb, nested = find_shape(slide, key)
    if shape is None:
        return
    zone = (tb[0] - 0.6, tb[1] - 0.25, tb[2] + 0.8, tb[3] + 0.45)
    for s, b, n in list(iter_shapes(slide.shapes)):
        if s is shape or n or s.shape_type in (6, 13):  # 组合、图片（logo）不动
            continue
        if s.has_text_frame and s.text_frame.text.strip():
            continue
        bw, bh = max(b[2], 0.01), max(b[3], 0.01)
        if bw * bh < 1.0 and overlap((b[0], b[1], bw, bh), zone) > 0.6 * bw * bh:
            remove(s)


def fill_toc(slide, items, slots, W, log, prof_slide=None, current=None):
    keys, nums, subs = slots.get("items", []), slots.get("numbers", []), slots.get("item_subs", [])
    prof_slide = prof_slide or {}
    i0 = prof_slide.get("toc_current")
    titles = [str((x[0] if x else "") if isinstance(x, list) else x) for x in items]
    if prof_slide.get("toc_mode") == "paragraphs":  # 一个文本框里一段一个议题
        nav.fill_paragraphs(slide, keys[0], titles, current, i0, prof_slide.get("toc_mark"), find_shape, log)
        return
    if i0 is not None:  # 模板高亮了某一项（议程式章节页）：高亮跟着当前章节走，没有当前章就都用普通样式
        n = min(len(titles), len(keys))
        for i in range(n, len(keys)):  # 先删多余条目（连同序号、说明、装饰），再填字
            drop_toc_entry(slide, keys[i], nums[i] if i < len(nums) else None, subs[i] if i < len(subs) else None)
        nav.fill_tracker(slide, {"items": keys[:n], "current": i0, "mark": None, "axis": "col"}, titles[:n], current,
                         find_shape, remove, log, label="议程")
        if nums:
            nav.fill_tracker(slide, {"items": nums[:n], "current": i0, "mark": None, "axis": "col"},
                             [f"{i + 1:02d}" for i in range(min(n, len(nums)))], current, find_shape, remove, log,
                             label="议程序号")
        for i in range(min(n, len(subs))):
            if subs[i]:
                v = items[i] if isinstance(items[i], list) else [items[i]]
                fill_slot(slide, subs[i], v[1] if len(v) > 1 and v[1] else None, W, log)
        return
    if len(items) > len(keys):
        log.append(f"目录有 {len(items)} 项，模板只有 {len(keys)} 个位置：多出的没放进去（请合并章节）")
    for i, key in enumerate(keys):
        sub = subs[i] if i < len(subs) else None
        if i < len(items):
            v = items[i] if isinstance(items[i], list) else [items[i]]
            if sub:  # 模板里副标题是单独的文本框
                fill_slot(slide, key, v[:1], W, log)
                fill_slot(slide, sub, v[1] if len(v) > 1 and v[1] else None, W, log)
            else:
                fill_slot(slide, key, [x for x in v if x], W, log)
            if i < len(nums):
                fill_slot(slide, nums[i], f"{i + 1:02d}", W, log)
        else:
            drop_toc_entry(slide, key, nums[i] if i < len(nums) else None, sub)


def drop_toc_entry(slide, key, num_key, sub_key=None):
    """删掉多余的目录条目：条目所在的组合整体删除；不在组合里就删条目、序号和压在它们下面的装饰"""
    shape, box, nested = find_shape(slide, key)
    if shape is None:
        return
    W = inch(slide.part.package.presentation_part.presentation.slide_width)
    H = inch(slide.part.package.presentation_part.presentation.slide_height)
    if nested:
        g = top_group(slide, shape)
        gb = next(b for s, b, _ in iter_shapes(slide.shapes) if s._element is g)
        if gb[2] * gb[3] < W * H * 0.2:
            g.getparent().remove(g)
            return
    zone = [box]
    for k in (num_key, sub_key):
        ns, nb, _ = find_shape(slide, k) if k else (None, None, None)
        if ns is not None:
            zone.append(nb)
            remove(ns)
    xs = [b[0] for b in zone] + [b[0] + b[2] for b in zone]
    ys = [b[1] for b in zone] + [b[1] + b[3] for b in zone]
    # 条目旁边的序号底、分隔线、竖条都算这一条的装饰（向外多扩一点）
    u = (min(xs) - 0.35, min(ys) - 0.15, max(xs) - min(xs) + 0.5, max(ys) - min(ys) + 0.4)
    remove(shape)
    for s, b, nested in list(iter_shapes(slide.shapes)):
        if nested or s.has_text_frame and s.text_frame.text.strip():
            continue
        bw, bh = max(b[2], 0.01), max(b[3], 0.01)
        if bw * bh < W * H * 0.08 and overlap((b[0], b[1], bw, bh), u) > 0.6 * bw * bh:
            remove(s)


def fix_paths(body, rel):
    """图片路径相对 deck.json 解析（不管在哪个目录运行 build）"""
    if not body:
        return
    if body.get("type") == "image" and body.get("path"):
        body["path"] = rel(body["path"])
        if not os.path.exists(body["path"]):
            sys.exit(f"找不到图片 {body['path']}")
    for k in ("left", "right"):
        if isinstance(body.get(k), dict) and body[k].get("type"):
            fix_paths(body[k], rel)


def body_area(spec, prof_slide, W, H):
    if spec.get("area"):
        return tuple(spec["area"])
    if prof_slide.get("area"):
        return tuple(prof_slide["area"])
    return (0.8, 1.6, W - 1.6, H - 2.2)


def build(deck_path, out=None):
    deck = load_json(deck_path)
    base = os.path.dirname(os.path.abspath(deck_path))
    rel = lambda p: p if os.path.isabs(p) else os.path.join(base, p)  # noqa: E731
    template = rel(deck["template"])
    prof_path = rel(deck.get("profile", "work/profile.json"))
    if not os.path.exists(prof_path):
        sys.exit(f"找不到 {prof_path}：先运行 inspect_template.py {template} --out {os.path.dirname(prof_path)}")
    prof = load_json(prof_path)
    tokens = {**prof["tokens"], **{k: v.lstrip("#") if isinstance(v, str) else v
                                   for k, v in (deck.get("tokens") or {}).items()}}
    prs = Presentation(template)
    W, H = inch(prs.slide_width), inch(prs.slide_height)
    src = list(prs.slides)
    n0 = len(src)
    report = {"slides": [], "warnings": []}
    for i, spec in enumerate(deck["slides"]):
        draw.WARN.clear()
        log = []
        fi = spec["from"]
        if not 0 <= fi < n0:
            sys.exit(f"第 {i + 1} 页 from={fi} 超出模板页数 {n0}")
        ps = prof["slides"][fi]
        slide = clone(prs, src[fi])
        if spec.get("delete_samples", True):
            for key in ps.get("samples", []):
                shape, _, _ = find_shape(slide, key)
                if shape is not None and key not in spec.get("keep", []):
                    remove(shape)
        for key in spec.get("delete", []):
            shape, _, _ = find_shape(slide, key)
            if shape is not None:
                remove(shape)
        apply_slots(slide, spec, ps, W, log)
        area = body_area(spec, ps, W, H)
        area0 = area  # 原始内容区（lint 检查新内容有没有画出界）
        before = set(slide.shapes._spTree)  # 持有节点引用：lxml 代理对象的 id() 不稳定
        fix_paths(spec.get("body"), rel)
        if spec.get("body") or spec.get("lead") or spec.get("conclusion"):
            bg = spec.get("bg") or ps.get("bg", "FFFFFF")
            T = Theme(tokens, bg.lstrip("#"))
            if spec.get("source"):
                area = recipes.source(slide, T, area, spec["source"])
            if spec.get("conclusion"):
                area = recipes.conclusion(slide, T, area, spec["conclusion"])
            if spec.get("lead"):
                area = recipes.lead(slide, T, area, spec["lead"])
            if spec.get("body"):
                recipes.draw(slide, T, area, spec["body"])
        for el in slide.shapes._spTree:  # 新画的形状加 xm 前缀，lint 据此区分「新内容」和「模板原有」
            if el not in before and el.tag != qn("p:extLst"):
                c = el.find(".//" + qn("p:cNvPr"))
                if c is not None:
                    c.set("name", "xm " + c.get("name", ""))
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]
        msgs = log + list(draw.WARN)
        report["slides"].append({"page": i + 1, "from": fi, "role": ps["role"], "area": list(area0),
                                 "body": (spec.get("body") or {}).get("type"), "messages": msgs})
        for m in msgs:
            report["warnings"].append(f"第 {i + 1} 页：{m}")
    drop_originals(prs, n0)
    out = out or rel(deck.get("output", "out/deck.pptx"))
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    report["output"] = os.path.basename(out)
    report["profile"] = os.path.relpath(prof_path, os.path.dirname(os.path.abspath(out)))
    save_json(out + ".build.json", report)
    print(f"已生成 {out}（{len(deck['slides'])} 页）")
    for w in report["warnings"]:
        print("  ⚠️", w)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("--out")
    a = ap.parse_args()
    build(a.deck, a.out)


if __name__ == "__main__":
    main()
